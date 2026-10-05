"""Use case: Evaluates the data partner on attributed public banking and retail data.

What it does: Retains verified source artifacts, exercises disposable real services,
and checks executed answers against independent standard-library Decimal calculations.
"""

import argparse
import csv
import hashlib
import io
import json
import os
from collections import Counter, defaultdict
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from datetime import datetime, timezone
from decimal import ROUND_HALF_EVEN, Decimal, localcontext
from fnmatch import fnmatch
from functools import partial
from pathlib import Path
from time import perf_counter
from typing import Any
from urllib.request import Request, urlopen
from uuid import uuid4
from zipfile import ZipFile

from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.engine import make_url

from execplus.bootstrap import build_runtime
from execplus.config import Settings
from execplus.infrastructure.identity import LocalSessionIdentity
from execplus.infrastructure.object_storage import S3ObjectStorage
from execplus.main import create_app

ROOT = Path(__file__).resolve().parents[1]
DIRECTORY = ROOT / "data" / "realdata-evaluation"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SOURCES = {
    "bank": {
        "title": "UCI Bank Marketing: publisher-provided bank.csv sample",
        "url": "https://archive.ics.uci.edu/static/public/222/bank+marketing.zip",
        "documentation": "https://archive.ics.uci.edu/dataset/222/bank+marketing",
        "attribution": "Moro, S., Rita, P., & Cortez, P. (2014). Bank Marketing.",
        "doi": "https://doi.org/10.24432/C5K306",
        "archive": "bank-marketing.zip",
        "license": "CC BY 4.0",
        "license_url": "https://creativecommons.org/licenses/by/4.0/",
    },
    "retail": {
        "title": "UCI Online Retail: first 10,000 original transaction rows",
        "url": "https://archive.ics.uci.edu/static/public/352/online+retail.zip",
        "documentation": "https://archive.ics.uci.edu/dataset/352/online+retail",
        "attribution": "Chen, D. (2015). Online Retail.",
        "doi": "https://doi.org/10.24432/C5BW33",
        "archive": "online-retail.zip",
        "license": "CC BY 4.0",
        "license_url": "https://creativecommons.org/licenses/by/4.0/",
    },
}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fetch(directory: Path, source: dict[str, str]) -> Path:
    path = directory / source["archive"]
    if not path.exists():
        request = Request(source["url"], headers={"User-Agent": "ExecPlus-public-evaluation/1.0"})
        with urlopen(request, timeout=120) as response:
            content = response.read(40 * 1024 * 1024 + 1)
        if len(content) > 40 * 1024 * 1024:
            raise ValueError("Public archive exceeds the bounded evaluation download")
        temporary = path.with_suffix(".partial")
        temporary.write_bytes(content)
        with ZipFile(temporary) as archive:
            if archive.testzip():
                raise ValueError("Public archive checksum validation failed")
        temporary.replace(path)
    return path


def prepare(directory: Path) -> list[dict[str, Any]]:
    directory.mkdir(parents=True, exist_ok=True)
    manifest_path = directory / "sources.json"
    if manifest_path.exists():
        manifest: list[dict[str, Any]] = json.loads(manifest_path.read_text())
        for entry in manifest:
            for artifact in entry["artifacts"]:
                if digest(directory / artifact["file"]) != artifact["sha256"]:
                    raise ValueError("Retained public artifact changed; use a fresh directory")
        return manifest
    bank = fetch(directory, SOURCES["bank"])
    with ZipFile(bank) as outer, ZipFile(io.BytesIO(outer.read("bank.zip"))) as inner:
        (directory / "bank-original.csv").write_bytes(inner.read("bank.csv"))
    with (directory / "bank-original.csv").open(newline="", encoding="utf-8") as stream:
        bank_rows = list(csv.reader(stream, delimiter=";"))
    with (directory / "bank-comma.csv").open("w", newline="", encoding="utf-8") as stream:
        csv.writer(stream).writerows(bank_rows)
    retail = fetch(directory, SOURCES["retail"])
    with ZipFile(retail) as archive:
        original = directory / "online-retail-original.xlsx"
        original.write_bytes(archive.read("Online Retail.xlsx"))
    workbook = load_workbook(original, read_only=True, data_only=False, keep_links=False)
    subset = Workbook(write_only=True)
    sheet = subset.create_sheet("Original rows 1-10000")
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    sheet.append(list(next(rows)))
    for _, row in zip(range(10_000), rows, strict=False):
        sheet.append(list(row))
    workbook.close()
    subset.save(directory / "retail-first10000.xlsx")
    entries = [
        (
            "bank",
            ["bank-marketing.zip", "bank-original.csv", "bank-comma.csv"],
            "Publisher's 4,521-row sample retained. Comma derivative changes only CSV delimiter.",
            "bank-original.csv",
        ),
        (
            "retail",
            ["online-retail.zip", "online-retail-original.xlsx", "retail-first10000.xlsx"],
            "First 10,000 data rows in original order; values unchanged. No deduplication, "
            "return exclusion, missing-value imputation or renaming. Original exceeds upload "
            "size/row limits. This ordered subset is not representative or a scale benchmark.",
            "retail-first10000.xlsx",
        ),
    ]
    manifest = [
        {
            "id": key,
            **SOURCES[key],
            "retrieved_at": datetime.now(timezone.utc).isoformat(),
            "preparation": preparation,
            "evaluation_file": evaluation_file,
            "artifacts": [{"file": name, "sha256": digest(directory / name)} for name in files],
        }
        for key, files, preparation, evaluation_file in entries
    ]
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n")
    return manifest


def reference_rows(path: Path) -> list[dict[str, str]]:
    if path.suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as stream:
            return list(
                csv.DictReader(stream, delimiter=";" if path.name == "bank-original.csv" else ",")
            )
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        iterator = workbook.worksheets[0].iter_rows(values_only=True)
        names = [str(value) for value in next(iterator)]
        return [
            dict(zip(names, ("" if value is None else str(value) for value in row), strict=True))
            for row in iterator
        ]
    finally:
        workbook.close()


def aggregate(values: list[str], method: str) -> Decimal:
    numbers = [Decimal(value) for value in values if value.strip()]
    with localcontext() as context:
        context.prec = 80
        if method == "count":
            return Decimal(len(numbers))
        if method == "sum":
            return sum(numbers, Decimal(0))
        if method == "avg":
            return (sum(numbers, Decimal(0)) / len(numbers)).quantize(
                Decimal("0.000000000001"), rounding=ROUND_HALF_EVEN
            )
        if method == "min":
            return min(numbers)
        if method == "max":
            return max(numbers)
    raise ValueError("Unsupported independent reference calculation")


@contextmanager
def services(live: bool) -> Iterator[tuple[TestClient, dict[str, str]]]:
    url = os.environ["EXECPLUS_TEST_DATABASE_URL"]
    schema = "realdata_" + uuid4().hex
    bucket = "realdata-" + uuid4().hex
    admin = create_engine(url)
    isolated = make_url(url).update_query_dict({"options": f"-csearch_path={schema}"})
    model_options: dict[str, Any] = (
        {}
        if live
        else {
            "llm_mode": "disabled",
            "llm_selection_base_url": "",
            "llm_selection_model": "",
        }
    )
    settings = Settings(
        environment="test",
        database_url=isolated.render_as_string(hide_password=False),
        object_store_endpoint=os.getenv(
            "EXECPLUS_TEST_OBJECT_STORE_ENDPOINT", "http://localhost:9000"
        ),
        object_store_bucket=bucket,
        object_store_access_key=os.getenv("EXECPLUS_TEST_OBJECT_STORE_ACCESS_KEY", "execplus"),
        object_store_secret_key=os.getenv("EXECPLUS_TEST_OBJECT_STORE_SECRET_KEY", "change-me"),
        email_mode="disabled",
        **model_options,
    )
    if live and settings.llm_mode == "disabled":
        raise ValueError(
            "Live evaluation needs a configured model; no scripted model is substituted"
        )
    runtime = build_runtime(settings)
    assert isinstance(runtime.service.storage, S3ObjectStorage)
    assert isinstance(runtime.identity, LocalSessionIdentity)
    storage = runtime.service.storage.client
    created_bucket = False
    try:
        with admin.begin() as connection:
            connection.execute(text(f'CREATE SCHEMA "{schema}"'))
        config = Config(str(ROOT / "alembic.ini"))
        config.set_main_option("script_location", str(ROOT / "migrations"))
        with runtime.engine.begin() as connection:
            config.attributes["connection"] = connection
            command.upgrade(config, "head")
        storage.create_bucket(Bucket=bucket)
        created_bucket = True
        token = runtime.identity.provision("public-evaluation@example.test")
        with TestClient(create_app(settings, runtime)) as client:
            yield client, {"Authorization": "Bearer " + token}
    finally:
        if created_bucket:
            paginator = storage.get_paginator("list_objects_v2")
            for page in paginator.paginate(Bucket=bucket):
                for item in page.get("Contents", []):
                    storage.delete_object(Bucket=bucket, Key=item["Key"])
            storage.delete_bucket(Bucket=bucket)
        runtime.engine.dispose()
        with admin.begin() as connection:
            connection.execute(text(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE'))
        admin.dispose()


class Evaluation:
    def __init__(
        self, client: TestClient, headers: dict[str, str], only: str | None = None
    ) -> None:
        self.client = client
        self.headers = headers
        self.results: list[dict[str, Any]] = []
        self.only = only

    def request(self, method: str, path: str, status: int = 200, **kwargs: Any) -> Any:
        response = self.client.request(method, path, headers=self.headers, **kwargs)
        if response.status_code != status:
            code = response.json().get("error", {}).get("code", "unknown")
            raise AssertionError(f"Expected HTTP {status}, got {response.status_code}: {code}")
        return response.json()

    def case(self, case_id: str, check: Callable[[], Any]) -> None:
        if self.only and not fnmatch(case_id, self.only):
            return
        start = perf_counter()
        try:
            details = check()
            result = {"id": case_id, "passed": True, "details": details}
        except Exception as error:
            result = {"id": case_id, "passed": False, "error_type": type(error).__name__}
            if isinstance(error, AssertionError):
                result["reason"] = str(error)[:250]
        result["latency_ms"] = round((perf_counter() - start) * 1000, 2)
        self.results.append(result)
        print(f"{'PASS' if result['passed'] else 'FAIL'} {case_id}", flush=True)

    def query(
        self,
        root: str,
        rows: list[dict[str, str]],
        metric: str,
        aggregation: str,
        group: str | None = None,
        filter_column: str | None = None,
        filter_value: str = "",
    ) -> dict[str, Any]:
        selected = [row for row in rows if not filter_column or row[filter_column] == filter_value]
        request: dict[str, Any] = {"metric": metric, "aggregation": aggregation}
        if group:
            request["group_by"] = [group]
        if filter_column:
            request["filters"] = [
                {"column": filter_column, "operator": "eq", "value": filter_value}
            ]
        answer = self.request("POST", root + "/query", json=request)
        grouped: dict[str | None, list[str]] = defaultdict(list)
        for row in selected:
            grouped[(row[group].strip() or None) if group else "__all"].append(row[metric])
        expected = {key: aggregate(values, aggregation) for key, values in grouped.items()}
        actual = {
            (str(row[0]) if row[0] is not None else None) if group else "__all": Decimal(
                str(row[-1])
            )
            for row in answer["rows"]
        }
        assert expected == actual, "Independent Decimal result mismatch"
        assert answer["records_analyzed"] == len(rows), "Source row count changed"
        workspace = root.split("/")[2]
        replay = self.request(
            "POST", f"/workspaces/{workspace}/queries/{answer['lineage']['query_id']}/replay"
        )
        assert replay["rows"] == answer["rows"], "Receipt replay differed"
        return {
            "groups_checked": len(expected),
            "independent_reference": "stdlib Decimal precision80",
            "replay": "matched",
        }

    def discovery(
        self,
        root: str,
        rows: list[dict[str, str]],
        expected_metrics: tuple[str, str],
        expected_dimension: str,
    ) -> dict[str, Any]:
        answer = self.request("GET", root + "/discovery")
        assert answer["version"] == "discovery-v1"
        assert answer["shape"]["rows"] == len(rows)
        assert answer["sources"] and answer["suggestions"]
        expected_findings = {
            ("metric", metric, aggregation)
            for metric in expected_metrics
            for aggregation in ("min", "max", "avg")
        } | {("distribution", expected_dimension, "count")}
        actual_findings = [
            (finding["kind"], finding["metric"], finding["aggregation"])
            for finding in answer["findings"]
        ]
        assert len(actual_findings) == 7 and set(actual_findings) == expected_findings, (
            "All seven expected public-data findings must complete without duplicates"
        )
        expected_quality = {
            "missing_values": sum(not value.strip() for row in rows for value in row.values()),
            "duplicate_rows": len(rows) - len({tuple(row.values()) for row in rows}),
        }
        for code, count in expected_quality.items():
            observed = next((item for item in answer["quality"] if item["code"] == code), None)
            if count:
                assert observed and observed["text"].startswith(str(count) + " "), (
                    "Data-quality finding does not match independent original-row checks"
                )
            else:
                assert observed is None, "A quality issue was reported without source evidence"
        checked = 0
        for finding in answer["findings"]:
            if finding["kind"] == "metric":
                actual = Decimal(str(finding["query"]["rows"][0][-1]))
                expected = aggregate(
                    [row[finding["metric"]] for row in rows], finding["aggregation"]
                )
                assert actual == expected, "Automatic metric did not match independent reference"
                checked += 1
            else:
                expected_counts = Counter(row[finding["metric"]].strip() or None for row in rows)
                actual_counts = {row[0]: row[-1] for row in finding["query"]["rows"]}
                assert dict(expected_counts) == actual_counts, "Category distribution mismatch"
            workspace = root.split("/")[2]
            replay = self.request(
                "POST",
                f"/workspaces/{workspace}/queries/{finding['query']['lineage']['query_id']}/replay",
            )
            assert replay["rows"] == finding["query"]["rows"], "Automatic finding replay mismatch"
        assert checked == 6, "All six numerical findings must be independently verified"
        return {
            "findings": len(answer["findings"]),
            "exact_metrics_checked": checked,
            "quality_flags": len(answer["quality"]),
            "finding_titles": [finding["title"] for finding in answer["findings"]],
            "quality": answer["quality"],
            "independent_quality_counts": expected_quality,
            "all_findings_replayed": True,
        }

    def guidance(self, root: str, name: str) -> dict[str, Any]:
        answer = self.request("POST", root + "/ask", json={"question": f"What does {name} mean?"})
        assert answer["kind"] == "overview" and name in answer["message"]
        assert answer["sources"]
        return {"kind": answer["kind"], "sources": "present"}

    def ambiguous(self, root: str) -> dict[str, Any]:
        body = self.request(
            "POST",
            root + "/ask",
            status=422,
            json={"question": "What does nonexistent_field mean?"},
        )
        assert body["error"]["code"] == "clarification_required"
        return {"clarification": "unknown field; no invented definition"}

    def live_total(self, root: str, rows: list[dict[str, str]], metric: str) -> dict[str, Any]:
        answer = self.request(
            "POST", root + "/ask", json={"question": f"What is the total {metric}?"}
        )
        assert Decimal(str(answer["value"])) == aggregate([row[metric] for row in rows], "sum")
        return {"model_route": answer["lineage"]["model_route"], "exact_reference": "matched"}

    def live_guidance(self, root: str, question: str, focus: str) -> dict[str, Any]:
        answer = self.request("POST", root + "/ask", json={"question": question})
        assert answer.get("kind") == "overview", "Help request did not return dataset guidance"
        assert answer["guidance"]["focus"] == focus, "Guidance did not address requested focus"
        assert answer["sources"] and answer["guidance"]["suggestions"]
        return {"model_route": answer["model_route"], "focus": focus, "message": answer["message"]}

    def live_followup(
        self,
        root: str,
        rows: list[dict[str, str]],
        metric: str,
        filter_column: str,
        filter_value: str,
        breakdown: str,
    ) -> dict[str, Any]:
        thread = self.request("POST", root + "/threads", status=201)["id"]
        workspace = root.split("/")[2]
        thread_root = f"/workspaces/{workspace}/threads/{thread}"

        def ask(question: str) -> Any:
            result = self.request(
                "POST",
                thread_root + "/ask",
                json={"question": question, "request_id": str(uuid4())},
            )
            assert result["turn"]["status"] == "complete", (
                f"{question}: {result['turn']['status']} ({result['turn']['kind']}): "
                f"{result['turn'].get('message', '')}"
            )
            return result["answer"]

        selected = [row for row in rows if row[filter_column] == filter_value]
        records = ask(f"Show all records where {filter_column} equals {filter_value}")
        assert records["matched_records"] == len(selected), "Record filter matched wrong number"
        assert 0 < len(records["rows"]) <= len(selected), (
            "Matching source records must return a nonempty, bounded result page"
        )
        index = records["columns"].index(filter_column)
        assert all(row[index] == filter_value for row in records["rows"]), "Wrong filtered records"
        total = ask(f"What is their total {metric}?")
        assert Decimal(str(total["value"])) == aggregate([row[metric] for row in selected], "sum")
        assert ask("thanks")["kind"] == "overview"
        grouped = ask(f"Break that total down by {breakdown}")
        expected: dict[str | None, list[str]] = defaultdict(list)
        for row in selected:
            expected[row[breakdown].strip() or None].append(row[metric])
        actual = {
            str(row[0]) if row[0] is not None else None: Decimal(str(row[-1]))
            for row in grouped["rows"]
        }
        assert actual == {key: aggregate(values, "sum") for key, values in expected.items()}, (
            "Follow-up lost its metric or filter"
        )
        return {
            "turns": 4,
            "filtered_records": len(selected),
            "returned_records": len(records["rows"]),
            "groups_checked": len(actual),
            "model_route": grouped["lineage"]["model_route"],
            "exact_reference": "matched",
        }


def evaluate(
    directory: Path,
    manifest: list[dict[str, Any]],
    live: bool,
    only: str | None = None,
) -> dict[str, Any]:
    with services(live) as (client, headers):
        evaluation = Evaluation(client, headers, only)
        wid = evaluation.request(
            "POST",
            "/workspaces",
            status=201,
            json={"name": "Public data evaluation", "seat_limit": 3},
        )["id"]
        for entry in manifest:
            key = entry["id"]
            path = directory / entry["evaluation_file"]
            rows = reference_rows(path)
            did = evaluation.request(
                "POST", f"/workspaces/{wid}/datasets", status=201, json={"name": entry["title"]}
            )["id"]
            response = client.post(
                f"/workspaces/{wid}/datasets/{did}/uploads",
                params={"filename": path.name},
                headers={**headers, "Content-Type": XLSX if path.suffix == ".xlsx" else "text/csv"},
                content=path.read_bytes(),
            )
            if response.status_code != 201:
                evaluation.results.append(
                    {
                        "id": key + ":intake",
                        "passed": False,
                        "http_status": response.status_code,
                        "code": response.json().get("error", {}).get("code"),
                    }
                )
                continue
            assert response.json()["row_count"] == len(rows)
            evaluation.results.append(
                {
                    "id": key + ":intake",
                    "passed": True,
                    "rows": len(rows),
                    "file_sha256": digest(path),
                }
            )
            root = f"/workspaces/{wid}/datasets/{did}/uploads/{response.json()['id']}"
            metric, secondary, group, selected = (
                ("balance", "age", "job", "retired")
                if key == "bank"
                else ("Quantity", "UnitPrice", "Country", "France")
            )
            evaluation.case(
                key + ":sum-and-replay", partial(evaluation.query, root, rows, metric, "sum")
            )
            evaluation.case(
                key + ":average-and-replay", partial(evaluation.query, root, rows, secondary, "avg")
            )
            evaluation.case(
                key + ":grouped-and-replay",
                partial(evaluation.query, root, rows, metric, "sum", group=group),
            )
            evaluation.case(
                key + ":filter-and-replay",
                partial(
                    evaluation.query,
                    root,
                    rows,
                    metric,
                    "sum",
                    filter_column=group,
                    filter_value=selected,
                ),
            )
            evaluation.case(
                key + ":automatic-discovery",
                partial(
                    evaluation.discovery,
                    root,
                    rows,
                    ("balance", "duration") if key == "bank" else ("Quantity", "UnitPrice"),
                    group,
                ),
            )
            evaluation.case(key + ":column-guidance", partial(evaluation.guidance, root, metric))
            evaluation.case(
                key + ":unknown-field-clarification", partial(evaluation.ambiguous, root)
            )
            if live:
                evaluation.case(
                    key + ":live-natural-language-total",
                    partial(evaluation.live_total, root, rows, metric),
                )
                evaluation.case(
                    key + ":live-filter-followup-social-breakdown",
                    partial(
                        evaluation.live_followup,
                        root,
                        rows,
                        metric,
                        group,
                        selected,
                        "education" if key == "bank" else "Description",
                    ),
                )
                evaluation.case(
                    key + ":live-roman-urdu-quality",
                    partial(
                        evaluation.live_guidance,
                        root,
                        "Mujhe is data ki quality ke bare mein batao",
                        "quality",
                    ),
                )
                evaluation.case(
                    key + ":live-urdu-next-steps",
                    partial(
                        evaluation.live_guidance,
                        root,
                        "اس ڈیٹا سے میں کیا سوال پوچھ سکتا ہوں؟",
                        "next_steps",
                    ),
                )
        return {
            "scope": "public-data local/test application evaluation",
            "production_evidence": False,
            "live_model": live,
            "case_filter": only,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "sources": manifest,
            "cases": evaluation.results,
            "passed": sum(case["passed"] for case in evaluation.results),
            "total": len(evaluation.results),
            "limitations": [
                "Bank uses publisher-provided sample; retail uses first10,000 rows, "
                "not full population.",
                "Exact calculations are independent Decimal checks, "
                "not subjective answer-quality labels.",
                "Text grouping follows the existing query contract: surrounding whitespace "
                "is trimmed and blank text becomes NULL; originals remain unchanged.",
                "This does not validate production load, forecasting, all formats, "
                "or every business interpretation.",
            ],
        }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--directory", type=Path, default=DIRECTORY)
    parser.add_argument("--download-only", action="store_true")
    parser.add_argument("--live-model", action="store_true")
    parser.add_argument("--only", help="Optional case-ID glob for a targeted rerun")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    manifest = prepare(args.directory)
    if args.download_only:
        print(f"Prepared {len(manifest)} attributed public datasets in {args.directory}")
        return
    output = args.output or args.directory / (
        "evaluation-" + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + ".json"
    )
    if output.exists():
        parser.error("Use a fresh output path; prior evaluations must be retained")
    started = perf_counter()
    report = evaluate(args.directory, manifest, args.live_model, args.only)
    report["elapsed_ms"] = round((perf_counter() - started) * 1000, 2)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, indent=2) + "\n")
    print(f"{report['passed']}/{report['total']} passed; evidence: {output}")
    if report["passed"] != report["total"]:
        parser.exit(1, "Public-data evaluation has failures; inspect the retained report.\n")


if __name__ == "__main__":
    main()
