"""Use case: Exercises exact comparisons through saved studies and protected outputs.

What it does: Uses synthetic source files to verify mapping, replay, policy races,
shared boards and independently parsed exports without model or provider calls.
"""

from copy import deepcopy
from uuid import UUID, uuid4

import pytest
from sqlalchemy import select, update

from execplus.infrastructure.persistence import schema as s
from test_analysis_access_policies import grant, revise
from test_artifacts import artifact_url
from test_result_exports import download_text
from test_studies import add_member
from test_understanding import confirm, context
from test_workspace_integration import dataset, identity, upload, workspace

BEFORE = b"product,revenue,quantity\nA,100,10\nB,200,10\nD,30,1\n"
AFTER = b"item,sales,units\nA,240,20\nB,125,5\nC,80,2\n"


def source(env, headers, wid, content, amount="revenue", quantity="quantity"):
    did = dataset(env, headers, wid)
    response = upload(env, headers, wid, did, content)
    assert response.status_code == 201, response.text
    uid = response.json()["id"]
    root = f"/workspaces/{wid}/datasets/{did}/uploads/{uid}"
    current = context(env, headers, root)
    current["definition"]["domain"] = "sales"
    for column in current["definition"]["columns"]:
        if column["name"] == amount:
            column.update(role="metric", unit="currency", currency="USD")
        elif column["name"] == quantity:
            column.update(role="metric", unit="items", currency="")
        else:
            column["role"] = "dimension"
    response = confirm(env, headers, root, current)
    assert response.status_code == 201, response.text
    return did, uid, root


def prepare_contribution(env, before=BEFORE, after=AFTER, decompose=True):
    owner, actor = identity(env, f"contribution-{uuid4().hex}@example.test")
    wid = workspace(env, owner)
    did, uid, root = source(env, owner, wid, before)
    other_did, other_uid, other_root = source(env, owner, wid, after, "sales", "units")
    method = contribution_method(env, owner, root, other_root, decompose)
    return owner, actor, wid, did, uid, root, other_did, other_uid, other_root, method


def contribution_method(env, headers, root, other_root, decompose=True):
    current = env.client.get(other_root + "/adaptive-views", headers=headers).json()
    right = other_root.split("/")
    baseline = dict(column="revenue", segment="product")
    comparison = dict(
        column="sales",
        segment="item",
        dataset_id=right[4],
        upload_id=right[6],
        revision_id=current["revision_id"],
        understanding_id=current["understanding_id"],
    )
    if decompose:
        baseline["quantity_column"] = "quantity"
        comparison["quantity_column"] = "units"
    return dict(
        kind="contribution",
        mode="dataset",
        baseline=baseline,
        comparison=comparison,
        additive_rows_confirmed=True,
        decompose=decompose,
    )


def create_contribution(env, headers, root, method):
    current = env.client.get(root + "/adaptive-views", headers=headers).json()
    return env.client.post(
        root + "/studies",
        headers=headers,
        json=dict(
            name="Observed sales change",
            question="Which products contributed to this change?",
            revision_id=current["revision_id"],
            understanding_id=current["understanding_id"],
            method=method,
        ),
    )


def test_two_source_study_exact_mapping_replay_and_immutable_evidence(integration):
    env = integration
    owner, _, wid, did, _, root, other_did, _, _, method = prepare_contribution(env)
    response = create_contribution(env, owner, root, method)
    assert response.status_code == 201, response.text
    body = response.json()
    value = body["contribution"]
    assert value["net_change"]["display"] == "115"
    assert value["reconciled"] and value["pvm"]["reconciled"]
    assert {k: v["display"] for k, v in value["pvm"]["totals"].items()} == dict(
        volume="75", mix="-75", unit_revenue="65", entries="80", exits="-30"
    )
    assert [r["segment"] for r in value["rows"]] == ["A", "C", "B", "D"]
    version = body["version"]
    assert [r["dataset_id"] for r in version["evidence"]["sources"]] == [did, other_did]
    assert body["study"]["shared"] is False
    assert body["results"][0]["columns"] == ["product", "revenue", "quantity"]
    assert body["results"][1]["columns"] == ["item", "sales", "units"]
    url = f"/workspaces/{wid}/study-versions/{version['id']}"
    assert env.client.get(url, headers=owner).json() == body
    for result in body["results"]:
        qid = result["lineage"]["query_id"]
        replay = env.client.post(f"/workspaces/{wid}/queries/{qid}/replay", headers=owner)
        assert replay.status_code == 200, replay.text
        assert replay.json()["rows"] == result["rows"]
        assert len(result["lineage"]["receipt"]["sources"]) == 1
        assert len(result["lineage"]["receipt"]["study_binding"]["sources"]) == 2


def test_period_windows_use_explicit_dates_and_report_missing_coverage(integration):
    env = integration
    owner, _ = identity(env, f"period-{uuid4().hex}@example.test")
    wid = workspace(env, owner)
    did, uid, root = source(
        env,
        owner,
        wid,
        b"date,product,revenue\n2026-01-01,A,10\n2026-01-02,B,\n"
        b"2026-02-01,A,15\n2026-02-02,C,4\n2026-03-01,A,999\n",
    )
    current = env.client.get(root + "/adaptive-views", headers=owner).json()
    assert current["date_columns"] == ["date"]
    left = dict(
        column="revenue",
        segment="product",
        date_column="date",
        start="2026-01-01",
        end="2026-01-31",
    )
    right = {
        **left,
        "start": "2026-02-01",
        "end": "2026-02-28",
        "dataset_id": did,
        "upload_id": uid,
        "revision_id": current["revision_id"],
        "understanding_id": current["understanding_id"],
    }
    response = create_contribution(
        env,
        owner,
        root,
        dict(
            kind="contribution",
            mode="period",
            baseline=left,
            comparison=right,
            additive_rows_confirmed=True,
        ),
    )
    assert response.status_code == 201, response.text
    value = response.json()["contribution"]
    assert value["baseline_rows"] == value["comparison_rows"] == 2
    assert value["excluded_missing"] == 1 and value["net_change"] is None
    assert value["comparison_total"]["display"] == "19"
    assert not value["reconciled"]
    original = response.json()
    new_uid = upload(
        env, owner, wid, did, b"date,product,revenue\n2026-01-01,A,10\n2026-02-01,A,25\n"
    ).json()["id"]
    new_root = root.rsplit("/", 1)[0] + "/" + new_uid
    meaning = context(env, owner, new_root)
    assert confirm(env, owner, new_root, meaning).status_code == 201
    current = env.client.get(new_root + "/adaptive-views", headers=owner).json()
    rerun = env.client.post(
        f"/workspaces/{wid}/studies/{original['study']['id']}/rerun",
        headers=owner,
        json=dict(
            upload_id=new_uid,
            parent_id=original["version"]["id"],
            revision_id=current["revision_id"],
            understanding_id=current["understanding_id"],
        ),
    )
    assert rerun.status_code == 201, rerun.text
    assert rerun.json()["contribution"]["net_change"]["display"] == "15"
    assert all(
        source["upload_id"] == new_uid for source in rerun.json()["version"]["evidence"]["sources"]
    )
    assert (
        rerun.json()["version"]["method"]["baseline"] == original["version"]["method"]["baseline"]
    )
    reopened = env.client.get(
        f"/workspaces/{wid}/study-versions/{original['version']['id']}", headers=owner
    )
    assert reopened.json()["contribution"] == value


@pytest.mark.parametrize("format", ["pdf", "xlsx", "csv", "png"])
def test_contribution_downloads_match_source_and_include_both_lineages(integration, format):
    env = integration
    owner, _, wid, _, _, root, _, _, _, method = prepare_contribution(env)
    response = create_contribution(env, owner, root, method)
    assert response.status_code == 201, response.text
    body = response.json()
    url = f"/workspaces/{wid}/study-versions/{body['version']['id']}/export.{format}"
    response = env.client.get(url, headers=owner)
    text = download_text(response, format)
    assert "115" in text and "240" in text and "125" in text
    if format in {"csv", "png"}:
        assert all(
            source["revision_id"] in text for source in body["version"]["evidence"]["sources"]
        )
        assert "segment-contribution-v1" in text
    else:
        assert "contribution" in text.lower()


def test_private_receipts_become_shared_only_with_all_source_access(integration):
    env = integration
    owner, owner_actor, wid, did, _, root, other_did, _, _, method = prepare_contribution(env)
    member, actor = add_member(env, owner, wid)
    created = create_contribution(env, owner, root, method)
    assert created.status_code == 201, created.text
    body = created.json()
    base = f"/workspaces/{wid}"
    version = base + f"/study-versions/{body['version']['id']}"
    paths = [version, version + "/export.xlsx"]
    for result in body["results"]:
        paths += [
            base + f"/queries/{result['lineage']['query_id']}",
            artifact_url(
                f"/workspaces/{wid}/datasets/{result['lineage']['dataset_id']}",
                "query_result",
                result["lineage"]["query_id"],
            ),
        ]
    for path in paths:
        assert env.client.get(path, headers=member).status_code == 404, path
    assert (
        env.client.patch(
            base + f"/studies/{body['study']['id']}", headers=owner, json={"shared": True}
        ).status_code
        == 204
    )
    for path in paths:
        assert env.client.get(path, headers=member).status_code == 200, path
    board = env.client.post(
        base + "/study-boards",
        headers=owner,
        json=dict(name="Observed contributions", shared=True, pins=[body["version"]["id"]]),
    )
    assert board.status_code == 201, board.text
    board_path = base + f"/study-boards/{board.json()['id']}"
    assert env.client.get(board_path, headers=member).status_code == 200
    grant(
        env,
        owner,
        owner_actor,
        wid,
        other_did,
        actor,
        ["item", "sales", "units"],
        {"column": "item", "operator": "eq", "value": "A"},
    )
    for path in [*paths, board_path, board_path + "/export.csv"]:
        denied = env.client.get(path, headers=member)
        assert denied.status_code in {403, 404}, path
        assert "rows" not in denied.json()
    assert env.client.get(base + f"/datasets/{did}/studies", headers=member).json() == []
    assert env.client.get(base + "/study-boards", headers=member).json() == []


def test_source_change_during_compute_never_publishes_study_or_bound_receipts(
    integration, monkeypatch
):
    env = integration
    owner, actor, wid, did, _, root, other_did, other_uid, other_root, method = (
        prepare_contribution(env)
    )
    executor = env.runtime.analytics.executor
    execute = executor.execute
    queries = []
    current = context(env, owner, other_root)
    current["definition"]["description"] = "Meaning changed during computation"

    async def change(plan, *args):
        result = await execute(plan, *args)
        queries.append(str(plan.query_id))
        if len(queries) == 2:
            env.runtime.understanding.save(
                actor,
                UUID(wid),
                UUID(other_did),
                UUID(other_uid),
                UUID(current["revision_id"]),
                current["version"],
                "confirmed",
                current["definition"],
            )
        return result

    monkeypatch.setattr(executor, "execute", change)
    response = create_contribution(env, owner, root, method)
    assert response.status_code == 409, response.text
    assert env.client.get(f"/workspaces/{wid}/datasets/{did}/studies", headers=owner).json() == []
    for qid in queries:
        assert (
            env.client.post(f"/workspaces/{wid}/queries/{qid}/replay", headers=owner).status_code
            == 404
        )


def test_tampered_contribution_is_recomputed_and_rejected(integration):
    env = integration
    owner, _, wid, _, _, root, _, _, _, method = prepare_contribution(env)
    response = create_contribution(env, owner, root, method)
    assert response.status_code == 201, response.text
    body = response.json()
    evidence = deepcopy(body["version"]["evidence"])
    evidence["calculation"]["net_change"]["numerator"] = "999"
    with env.engine.begin() as connection:
        connection.execute(
            update(s.study_versions)
            .where(s.study_versions.c.id == UUID(body["version"]["id"]))
            .values(evidence=evidence)
        )
    url = f"/workspaces/{wid}/study-versions/{body['version']['id']}"
    for path in (url, url + "/export.xlsx"):
        response = env.client.get(path, headers=owner)
        assert (
            response.status_code == 409 and response.json()["error"]["code"] == "lineage_mismatch"
        )


def test_comparison_unit_mismatch_refuses_before_processing(integration):
    env = integration
    owner, _, _, _, _, root, _, _, other_root, method = prepare_contribution(env)
    current = context(env, owner, other_root)
    next(c for c in current["definition"]["columns"] if c["name"] == "sales")["currency"] = "EUR"
    assert confirm(env, owner, other_root, current).status_code == 201
    method = contribution_method(env, owner, root, other_root)
    response = create_contribution(env, owner, root, method)
    assert response.status_code == 422 and response.json()["error"]["code"] == "contribution_unit"
    with env.engine.connect() as connection:
        assert not connection.execute(select(s.query_executions)).all()


def test_both_sources_apply_rows_and_columns_before_any_contribution(integration):
    env = integration
    fixture = prepare_contribution(env)
    owner, owner_actor, wid, did, _, root, other_did, _, other_root, _ = fixture
    member, actor = add_member(env, owner, wid)
    grant(
        env,
        owner,
        owner_actor,
        wid,
        did,
        actor,
        ["product", "revenue"],
        dict(column="product", operator="eq", value="A"),
    )
    policy, body = grant(
        env,
        owner,
        owner_actor,
        wid,
        other_did,
        actor,
        ["item", "sales"],
        dict(column="item", operator="eq", value="A"),
    )
    method = contribution_method(env, member, root, other_root, False)
    response = create_contribution(env, member, root, method)
    assert response.status_code == 201, response.text
    value = response.json()
    assert value["contribution"]["net_change"]["display"] == "140"
    assert [row["segment"] for row in value["contribution"]["rows"]] == ["A"]
    assert all(result["records_analyzed"] == 1 for result in value["results"])
    assert "quantity" not in value["results"][0]["columns"]
    assert value["version"]["evidence"]["preparation"] == []
    assert "__restricted__" in response.text or "output_checksum" not in response.text
    url = f"/workspaces/{wid}/study-versions/{value['version']['id']}"
    assert env.client.get(url + "/export.xlsx", headers=member).status_code == 200
    assert env.client.get(url, headers=owner).status_code == 404
    revise(env, owner, policy, body)
    assert env.client.get(url, headers=member).status_code == 404
    assert env.client.get(f"/workspaces/{wid}/datasets/{did}/studies", headers=member).json() == []


def test_shared_projection_replay_rechecks_unshare_after_computation(integration, monkeypatch):
    env = integration
    owner, actor, wid, _, _, root, _, _, _, method = prepare_contribution(env)
    member, _ = add_member(env, owner, wid)
    response = create_contribution(env, owner, root, method)
    assert response.status_code == 201, response.text
    value = response.json()
    sid = UUID(value["study"]["id"])
    env.runtime.studies.share(actor, UUID(wid), sid, True)
    execute = env.runtime.analytics.executor.execute

    async def unshare(*args):
        result = await execute(*args)
        env.runtime.studies.share(actor, UUID(wid), sid, False)
        return result

    monkeypatch.setattr(env.runtime.analytics.executor, "execute", unshare)
    qid = value["results"][0]["lineage"]["query_id"]
    response = env.client.post(f"/workspaces/{wid}/queries/{qid}/replay", headers=member)
    assert response.status_code == 404, response.text


def test_secondary_policy_change_during_export_prevents_download(integration, monkeypatch):
    env = integration
    owner, owner_actor, wid, _, _, root, other_did, _, _, method = prepare_contribution(env)
    member, _actor = add_member(env, owner, wid)
    response = create_contribution(env, owner, root, method)
    assert response.status_code == 201, response.text
    value = response.json()
    env.runtime.studies.share(owner_actor, UUID(wid), UUID(value["study"]["id"]), True)
    renderer = env.runtime.result_exports.renderer
    original = renderer.render

    def revoke(*args):
        rendered = original(*args)
        # A direct service call avoids nesting a test HTTP client in an async route.
        env.runtime.access_policies.save(
            owner_actor,
            UUID(wid),
            UUID(other_did),
            0,
            True,
            [dict(user_id=str(owner_actor.id), columns=None, conditions=[])],
        )
        return rendered

    monkeypatch.setattr(renderer, "render", revoke)
    path = f"/workspaces/{wid}/study-versions/{value['version']['id']}/export.xlsx"
    assert env.client.get(path, headers=member).status_code == 404
    with env.engine.connect() as connection:
        assert not connection.execute(
            select(s.audit_events).where(s.audit_events.c.action == "result.exported")
        ).all()


def test_source_ids_cannot_cross_workspace_or_use_stale_meaning(integration):
    env = integration
    owner, _, wid, _, _, root, _, _, other_root, method = prepare_contribution(env)
    changed = context(env, owner, other_root)
    changed["definition"]["description"] = "Confirmed again"
    assert confirm(env, owner, other_root, changed).status_code == 201
    response = create_contribution(env, owner, root, method)
    assert response.status_code == 409, response.text
    other_wid = workspace(env, owner)
    _, _, other_root = source(env, owner, other_wid, AFTER, "sales", "units")
    method = contribution_method(env, owner, root, other_root)
    response = create_contribution(env, owner, root, method)
    assert response.status_code == 404, response.text
    assert (
        env.client.get(
            f"/workspaces/{wid}/datasets/{root.split('/')[4]}/studies", headers=owner
        ).json()
        == []
    )


def test_dataset_rerun_retains_fixed_comparison_and_original_version(integration):
    env = integration
    owner, _, wid, did, _, root, _, _, _, method = prepare_contribution(env)
    response = create_contribution(env, owner, root, method)
    assert response.status_code == 201, response.text
    first = response.json()
    new_uid = upload(
        env, owner, wid, did, b"product,revenue,quantity\nA,110,10\nB,200,10\nD,30,1\n"
    ).json()["id"]
    new_root = root.rsplit("/", 1)[0] + "/" + new_uid
    current = context(env, owner, new_root)
    assert confirm(env, owner, new_root, current).status_code == 201
    current = env.client.get(new_root + "/adaptive-views", headers=owner).json()
    response = env.client.post(
        f"/workspaces/{wid}/studies/{first['study']['id']}/rerun",
        headers=owner,
        json=dict(
            upload_id=new_uid,
            parent_id=first["version"]["id"],
            revision_id=current["revision_id"],
            understanding_id=current["understanding_id"],
        ),
    )
    assert response.status_code == 201, response.text
    second = response.json()
    assert second["version"]["number"] == 2
    assert second["contribution"]["net_change"]["display"] == "105"
    assert second["version"]["method"]["comparison"] == first["version"]["method"]["comparison"]
    reopened = env.client.get(
        f"/workspaces/{wid}/study-versions/{first['version']['id']}", headers=owner
    )
    assert reopened.status_code == 200 and reopened.json()["contribution"] == first["contribution"]


def test_larger_allowed_comparison_keeps_all_records_in_excel(integration):
    from io import BytesIO

    from openpyxl import load_workbook

    env = integration
    before = (
        "product,revenue,quantity\n" + "".join(f"P{i % 20},1,1\n" for i in range(1200))
    ).encode()
    after = ("item,sales,units\n" + "".join(f"P{i % 20},2,1\n" for i in range(1200))).encode()
    owner, _, wid, _, _, root, _, _, _, method = prepare_contribution(env, before, after)
    response = create_contribution(env, owner, root, method)
    assert response.status_code == 201, response.text
    value = response.json()
    assert value["contribution"]["net_change"]["display"] == "1200"
    assert [len(result["rows"]) for result in value["results"]] == [1200, 1200]
    response = env.client.get(
        f"/workspaces/{wid}/study-versions/{value['version']['id']}/export.xlsx", headers=owner
    )
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    raw_tables = [sheet for sheet in workbook if sheet.max_row == 1201 and sheet.max_column == 3]
    assert len(raw_tables) == 2
    for sheet in raw_tables:
        assert sum(str(row[0]).startswith("P") for row in list(sheet.values)[1:]) == 1200


def test_empty_filtered_populations_are_not_described_as_observed_performance(integration):
    env = integration
    owner, _, _, _, _, root, _, _, _, method = prepare_contribution(env)
    method["baseline"]["filters"] = [dict(column="product", operator="eq", value="No match")]
    method["comparison"]["filters"] = [dict(column="item", operator="eq", value="No match")]
    response = create_contribution(env, owner, root, method)
    assert response.status_code == 201, response.text
    value = response.json()
    assert value["contribution"]["rows"] == []
    assert value["contribution"]["net_change"]["display"] == "0"
    assert value["contribution"]["pvm"]["state"] == "no_matched_products"
    assert value["sample_count"] == 0
    assert any(
        "empty sums do not establish observed performance" in s for s in value["limitations"]
    )


def test_projection_dashboard_export_rechecks_sharing_at_final_download(integration, monkeypatch):
    env = integration
    owner, actor, wid, _, _, root, _, _, _, method = prepare_contribution(env)
    member, _ = add_member(env, owner, wid)
    response = create_contribution(env, owner, root, method)
    assert response.status_code == 201, response.text
    value = response.json()
    sid = UUID(value["study"]["id"])
    env.runtime.studies.share(actor, UUID(wid), sid, True)
    exports = env.runtime.result_exports
    original = exports._with_sources
    prepared = 0

    def unshare_after_final_replay(*args):
        nonlocal prepared
        result = original(*args)
        prepared += 1
        if prepared == 2:
            env.runtime.studies.share(actor, UUID(wid), sid, False)
        return result

    monkeypatch.setattr(exports, "_with_sources", unshare_after_final_replay)
    response = env.client.post(
        f"/workspaces/{wid}/exports/dashboard",
        headers=member,
        json=dict(
            name="Selected contribution records",
            format="xlsx",
            query_ids=[item["lineage"]["query_id"] for item in value["results"]],
        ),
    )
    assert prepared == 2
    assert response.status_code == 404, response.text[:120]
    with env.engine.connect() as connection:
        assert not connection.execute(
            select(s.audit_events).where(s.audit_events.c.action == "result.exported")
        ).all()


def test_ranked_fields_add_root_cause_evidence_that_replays_exactly(integration):
    env = integration
    owner, _ = identity(env, f"root-cause-{uuid4().hex}@example.test")
    wid = workspace(env, owner)
    _, uid, root = source(
        env,
        owner,
        wid,
        b"date,product,region,revenue\n2026-01-01,A,N,10\n2026-01-02,A,S,10\n"
        b"2026-01-03,B,N,5\n2026-02-01,A,N,30\n2026-02-02,A,S,10\n2026-02-03,B,N,4\n",
    )
    current = env.client.get(root + "/adaptive-views", headers=owner).json()
    left = dict(
        column="revenue",
        segment="product",
        dimensions=["region"],
        date_column="date",
        start="2026-01-01",
        end="2026-01-31",
    )
    right = {
        **left,
        "start": "2026-02-01",
        "end": "2026-02-28",
        "dataset_id": root.split("/")[4],
        "upload_id": uid,
        "revision_id": current["revision_id"],
        "understanding_id": current["understanding_id"],
    }
    method = dict(
        kind="contribution",
        mode="period",
        baseline=left,
        comparison=right,
        additive_rows_confirmed=True,
    )
    response = create_contribution(env, owner, root, method)
    assert response.status_code == 201, response.text
    body = response.json()
    cause = body["contribution"]["root_cause"]
    assert [f["field"] for f in cause["fields"]] == ["region", "product"]
    assert cause["fields"][0]["concentration"]["display"] == "1"
    assert cause["fields"][1]["concentration"]["numerator"] == "20"
    assert cause["fields"][1]["concentration"]["denominator"] == "21"
    drill = cause["drilldown"]
    assert drill["within_member"] == "N" and drill["reconciled"]
    assert {r["member"]: r["delta"]["display"] for r in drill["rows"]} == {"A": "20", "B": "-1"}
    assert body["version"]["evidence"]["method_version"] == "segment-contribution-v1"
    url = f"/workspaces/{wid}/study-versions/{body['version']['id']}"
    assert env.client.get(url, headers=owner).json() == body
    method["comparison"] = {**right, "dimensions": ["product"]}
    method["baseline"] = {**left, "dimensions": ["product"]}
    assert create_contribution(env, owner, root, method).status_code == 422
