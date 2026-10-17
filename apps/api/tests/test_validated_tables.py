"""Use case: Guards the single-traversal source optimization.

What it does: Independently checks exact CSV/XLSX values, retained rejection rules,
legacy parser compatibility and isolation when a parser serves concurrent requests.
"""

from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import Workbook

from execplus.application.validated_tables import read_validated_table
from execplus.domain.ingestion import IngestionError
from execplus.infrastructure import file_parser
from execplus.infrastructure.file_parser import StructuredFileParser
from test_file_parser import XLSX, workbook_bytes


@pytest.mark.parametrize(
    "content,format",
    [
        (b'item,amount,note\n A ,0.1234567890123456789,"line\none"\nB,-0.5,\n', "csv"),
        (b"item;amount\nA;0.1\nB;0\n", "csv;s"),
        (b"item\tamount\nA\t0.1\nB\t0\n", "csv;t"),
    ],
)
def test_combined_csv_preserves_original_text_nulls_delimiters_and_stream(content, format):
    parser = StructuredFileParser()
    stream = BytesIO(content)
    structure = parser.parse(stream, "fictional.csv", "text/csv", stored_format=format)
    expected = parser.read_table(stream, format)
    combined, table = parser.parse_table(stream, "fictional.csv", "text/csv", stored_format=format)
    assert combined == structure and table == expected
    assert not stream.closed and stream.tell() == 0


def test_combined_excel_keeps_dates_booleans_empty_values_and_exact_text(monkeypatch):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["day", "active", "amount", "empty"])
    sheet.append([date(2026, 1, 1), False, "9007199254740993.000000000000000001", None])
    sheet.append([datetime(2026, 1, 2, 15, 30), True, -0.125, ""])
    data = BytesIO()
    workbook.save(data)
    workbook.close()
    parser = StructuredFileParser()
    parser.parse(data, "fictional.xlsx", XLSX)
    expected = parser.read_table(data, "xlsx")
    original = file_parser.load_workbook
    calls = []

    def observe(*args, **kwargs):
        calls.append(1)
        return original(*args, **kwargs)

    monkeypatch.setattr(file_parser, "load_workbook", observe)
    structure, actual = parser.parse_table(data, "fictional.xlsx", XLSX)
    assert structure.row_count == 2 and actual == expected and len(calls) == 1
    assert actual.rows[0] == ("2026-01-01", "False", "9007199254740993.000000000000000001", "")
    assert actual.rows[1] == ("2026-01-02T15:30:00", "True", "-0.125", "")


@pytest.mark.parametrize(
    "content,name,mime,code",
    [
        (b"a,b\n1,2,3\n", "a.csv", "text/csv", "malformed_csv"),
        (b"a,a\n1,2\n", "a.csv", "text/csv", "invalid_header"),
        (b"a,b\n", "a.csv", "text/csv", "empty_table"),
        (b"a,b\n=1+1,2\n", "a.csv", "text/csv", "unsafe_formula"),
        (b"a,b\n1,\x00\n", "a.csv", "text/csv", "unsafe_content"),
        (b"<html>bad</html>", "a.csv", "text/csv", "content_mismatch"),
        (b"a,b\n1,2\n", "../a.csv", "text/csv", "unsafe_filename"),
        (workbook_bytes(multiple=True), "a.xlsx", XLSX, "multiple_sheets"),
        (workbook_bytes(formula=True), "a.xlsx", XLSX, "unsafe_formula"),
        (workbook_bytes(merged=True), "a.xlsx", XLSX, "merged_cells"),
        (b"PK\x03\x04bad", "a.xlsx", XLSX, "malformed_excel"),
    ],
)
def test_combined_path_retains_validation_rejections(content, name, mime, code):
    with pytest.raises(IngestionError) as failure:
        StructuredFileParser().parse_table(BytesIO(content), name, mime)
    assert failure.value.code == code


def test_shared_parser_has_no_cross_request_rows_or_retained_values():
    parser = StructuredFileParser()

    def read(index):
        raw = f"item,amount\nprivate-{index},{index}\n".encode()
        return read_validated_table(parser, BytesIO(raw), "a.csv", "text/csv", "csv")

    with ThreadPoolExecutor(max_workers=4) as pool:
        values = list(pool.map(read, range(20)))
    assert [table.rows for table in values] == [((f"private-{i}", str(i)),) for i in range(20)]
    assert vars(parser) == {}


def test_old_parser_port_falls_back_without_requiring_the_new_capability():
    parser = StructuredFileParser()

    class Legacy:
        parse = parser.parse
        read_table = parser.read_table

    table = read_validated_table(Legacy(), BytesIO(b"a,b\n1,2\n"), "a.csv", "text/csv", "csv")
    assert table.headers == ("a", "b") and table.rows == (("1", "2"),)


def test_stored_excel_format_cannot_change_during_combined_read():
    with pytest.raises(IngestionError) as failure:
        read_validated_table(
            StructuredFileParser(), BytesIO(workbook_bytes()), "a.xlsx", XLSX, "csv"
        )
    assert failure.value.code == "unsupported_format"
