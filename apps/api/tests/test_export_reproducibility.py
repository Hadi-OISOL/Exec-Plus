"""Use case: Checks reproducible, readable exports without weakening exact evidence.

What it does: Opens real files to verify stable bytes, safe distinct worksheet names,
visible metadata, timezone handling and unchanged exact typed cells.
"""

import json
import time
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from execplus.domain.ingestion import IngestionError
from execplus.domain.pdf_exports import PdfSection
from execplus.infrastructure.exports.portable import PortableRenderer
from test_portable_exports import content


def test_identical_fixed_report_has_identical_bytes_across_wall_clock_changes():
    renderer = PortableRenderer()
    report = content()
    first = {kind: renderer.render(report, kind) for kind in ("pdf", "xlsx", "csv")}
    time.sleep(2.05)  # Cross the ZIP format's two-second timestamp resolution.
    second = {kind: renderer.render(report, kind) for kind in first}
    assert [kind for kind in first if first[kind] != second[kind]] == []


def test_named_sheets_include_visible_summary_method_and_exact_evidence():
    report = content()
    report = replace(
        report,
        document=replace(
            report.document,
            sections=(
                *report.document.sections,
                PdfSection("Methodology", ("Sum of recorded amounts; no estimated values.",)),
            ),
        ),
    )
    workbook = load_workbook(BytesIO(PortableRenderer().render(report, "xlsx")))
    assert workbook.sheetnames == [
        "Summary",
        "Original result",
        "Methodology",
        "Evidence",
        "Read me",
    ]
    assert all(sheet.sheet_state == "visible" for sheet in workbook)
    summary = str(list(workbook["Summary"].values))
    assert report.document.title in summary and "2026-10-16" in summary
    method = str(list(workbook["Methodology"].values))
    assert "Sum of recorded amounts; no estimated values." in method
    assert "Source checksum: fictional-source" in method
    evidence = json.loads("".join(row[1] for row in list(workbook["Evidence"].values)[1:]))
    assert evidence["tables"][0]["rows"][0][0]["value"] == "9007199254740993.000000000000000001"
    assert workbook.properties.created == datetime(2026, 10, 16)
    assert workbook.properties.modified == datetime(2026, 10, 16)


def test_worksheet_names_are_bounded_distinct_and_keep_original_name_mapping():
    report = content((("=formula remains literal",),))
    names = ("Summary", "summary", "History", "'Revenue:/North?*[]\\'", "Long " * 20)
    report = replace(report, tables=tuple(replace(report.tables[0], name=name) for name in names))
    workbook = load_workbook(BytesIO(PortableRenderer().render(report, "xlsx")))
    assert len({name.casefold() for name in workbook.sheetnames}) == len(workbook.sheetnames)
    assert all(0 < len(name.encode("utf-16-le")) // 2 <= 31 for name in workbook.sheetnames)
    assert all(not any(char in name for char in "[]:*?/\\") for name in workbook.sheetnames)
    assert "History" not in workbook.sheetnames
    mapping = list(workbook["Read me"].values)
    for original in names:
        assert any(original in row for row in mapping)


def test_pdf_utc_label_matches_actual_offset_timestamp():
    report = content()
    report = replace(
        report,
        document=replace(
            report.document,
            generated_at=datetime(2026, 10, 17, 2, 30, tzinfo=timezone(timedelta(hours=5))),
        ),
    )
    reader = PdfReader(BytesIO(PortableRenderer().render(report, "pdf")))
    assert "Generated 2026-10-16 21:30 UTC" in reader.pages[0].extract_text()


@pytest.mark.parametrize("kind", ["pdf", "xlsx", "csv"])
def test_ambiguous_naive_report_time_is_refused(kind):
    report = content()
    report = replace(report, document=replace(report.document, generated_at=datetime(2026, 10, 17)))
    with pytest.raises(IngestionError, match="time zone"):
        PortableRenderer().render(report, kind)


def test_excel_largest_allowed_cell_count_retains_rows_types_and_formula_defenses():
    from datetime import date
    from decimal import Decimal

    from execplus.domain.exports import MAX_EXPORT_CELLS, validate_export

    # 19,999 rows plus headers, ten columns: the actual 200,000-cell bound.
    rows = tuple(
        (
            index,
            Decimal("0.100000000000000001"),
            date(2026, 1, 1),
            None,
            "",
            "=1+1",
            "USD",
            True,
            "A",
            -1,
        )
        for index in range(19_999)
    )
    report = content(rows)
    assert (len(rows) + 1) * len(rows[0]) == MAX_EXPORT_CELLS
    result = PortableRenderer().render(report, "xlsx")
    workbook = load_workbook(BytesIO(result), read_only=True)
    selected = workbook["Original result"]
    iterator = selected.iter_rows()
    next(iterator)
    seen = 0
    for row in iterator:
        assert row[0].value == str(seen)
        assert row[1].value == "0.100000000000000001"
        assert row[2].value == "2026-01-01"
        assert row[5].value == "=1+1" and row[5].data_type == "s"
        seen += 1
    assert seen == len(rows)
    evidence = json.loads("".join(row[1] for row in list(workbook["Evidence"].values)[1:]))
    last = evidence["tables"][0]["rows"][-1]
    assert last[0] == {"type": "integer", "value": "19998"}
    assert last[3] == {"type": "null", "value": ""}
    assert last[4] == {"type": "string", "value": ""}
    workbook.close()
    with pytest.raises(IngestionError, match="200,000-cell"):
        validate_export(replace(report, tables=(replace(report.tables[0], rows=(*rows, rows[0])),)))


def test_excel_long_methodology_preserves_all_text_in_order_without_tempfiles():
    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    before = set(ALL_TEMP_FILES)
    report = content()
    paragraph = "Confirmed source definition: " + "field\U0001f331" * 9000
    report = replace(
        report,
        document=replace(report.document, sections=(PdfSection("Definition", (paragraph,)),)),
    )
    workbook = load_workbook(BytesIO(PortableRenderer().render(report, "xlsx")))
    rows = list(workbook["Methodology"].values)[1:]
    assert "".join(row[1] for row in rows) == paragraph
    assert all(len(row[1].encode("utf-16-le")) // 2 <= 32767 for row in rows)
    assert [row[2] for row in rows] == [f"1 / {i} of {len(rows)}" for i in range(1, len(rows) + 1)]
    assert set(ALL_TEMP_FILES) == before
    workbook.close()


@pytest.mark.parametrize("field", ["title", "section", "paragraph"])
def test_excel_rejects_invalid_metadata_before_creating_temporary_files(field):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    before = set(ALL_TEMP_FILES)
    report = content()
    document = report.document
    if field == "title":
        document = replace(document, title="PRIVATE\x00TITLE")
    else:
        document = replace(
            document,
            sections=(
                PdfSection(
                    "PRIVATE\x00TITLE" if field == "section" else "Source", ("PRIVATE\x00CONTENT",)
                ),
            ),
        )
    with pytest.raises(IngestionError) as error:
        PortableRenderer().render(replace(report, document=document), "xlsx")
    assert error.value.code == "export_text" and "PRIVATE" not in str(error.value)
    assert set(ALL_TEMP_FILES) == before


def test_excel_writer_failure_removes_its_temporary_source_files(monkeypatch):
    from pathlib import Path

    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    from execplus.infrastructure.exports import portable

    before = set(ALL_TEMP_FILES)
    created = []

    def failed_save(writer):
        created.extend(set(ALL_TEMP_FILES) - before)
        assert created and all(Path(path).exists() for path in created)
        raise OSError("Synthetic archive failure")

    monkeypatch.setattr(portable.ExcelWriter, "save", failed_save)
    with pytest.raises(OSError, match="Synthetic archive failure"):
        PortableRenderer().render(content(), "xlsx")
    assert set(ALL_TEMP_FILES) == before
    assert all(not Path(path).exists() for path in created)
