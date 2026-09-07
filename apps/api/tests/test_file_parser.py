"""Use case: Proves deterministic upload acceptance and rejection.

What it does: Generates synthetic CSV/XLSX fixtures without committing data or binary workbooks.
"""

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from execplus.domain.ingestion import IngestionError
from execplus.infrastructure.file_parser import MAX_BYTES, StructuredFileParser, validate_filename

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook_bytes(*, multiple=False, merged=False, formula=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["item", "amount"])
    sheet.append(["widget", "=1+1" if formula else 12])
    if multiple:
        workbook.create_sheet("extra")
    if merged:
        sheet.merge_cells("A1:B1")
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.mark.parametrize(
    "data,name,mime,expected",
    [
        (b"item,amount\nwidget,12\n", "sales.csv", "text/csv", ("csv", 1, 2, 0)),
        (b'item,note\nwidget,"two\nlines"\n', "SALES.CSV", "text/csv", ("csv", 1, 2, 0)),
        (b"\xef\xbb\xbfitem,amount\nwidget,-12.5\n", "sales.csv", "text/csv", ("csv", 1, 2, 0)),
        (workbook_bytes(), "sales.xlsx", XLSX, ("xlsx", 1, 2, 1)),
    ],
)
def test_valid_files(data, name, mime, expected):
    content = BytesIO(data)
    result = StructuredFileParser().parse(content, name, mime)
    assert (result.format, result.row_count, result.column_count, result.sheet_count) == expected
    assert not content.closed
    assert content.tell() == 0


@pytest.mark.parametrize(
    "data,name,mime,code",
    [
        (b"a,b\n1,2,3\n", "a.csv", "text/csv", "malformed_csv"),
        (b'a,b\n1,"unclosed', "a.csv", "text/csv", "malformed_csv"),
        (b"a,b\n\xff,2", "a.csv", "text/csv", "malformed_csv"),
        (b"a,a\n1,2", "a.csv", "text/csv", "invalid_header"),
        (b",b\n1,2", "a.csv", "text/csv", "invalid_header"),
        (b"a,b\n", "a.csv", "text/csv", "empty_table"),
        (b"", "a.csv", "text/csv", "empty_file"),
        (b"a,b\n1,2\n\n", "a.csv", "text/csv", "malformed_csv"),
        (b"a,b\n=1+1,2", "a.csv", "text/csv", "unsafe_formula"),
        (b"a,b\n @SUM(A1),2", "a.csv", "text/csv", "unsafe_formula"),
        (b"a,b\n\x00,2", "a.csv", "text/csv", "unsafe_content"),
        (b"MZfake", "a.csv", "text/csv", "content_mismatch"),
        (workbook_bytes(), "a.csv", "text/csv", "content_mismatch"),
        (b"a,b\n1,2", "a.xls", "application/vnd.ms-excel", "unsupported_format"),
        (b"a,b\n1,2", "a.exe", "text/csv", "unsupported_format"),
        (b"a,b\n1,2", "a.csv", "application/pdf", "unsupported_format"),
        (b"not excel", "a.xlsx", XLSX, "malformed_excel"),
        (b"PK\x03\x04bad", "a.xlsx", XLSX, "malformed_excel"),
        (workbook_bytes(multiple=True), "a.xlsx", XLSX, "multiple_sheets"),
        (workbook_bytes(merged=True), "a.xlsx", XLSX, "merged_cells"),
        (workbook_bytes(formula=True), "a.xlsx", XLSX, "unsafe_formula"),
        (b"a,b\n1,2", "../a.csv", "text/csv", "unsafe_filename"),
    ],
)
def test_rejections_have_stable_codes(data, name, mime, code):
    with pytest.raises(IngestionError) as error:
        StructuredFileParser().parse(BytesIO(data), name, mime)
    assert error.value.code == code
    assert len(str(error.value)) > 12


@pytest.mark.parametrize(
    "name",
    [
        "../a.csv",
        "/tmp/a.csv",
        "a\\b.csv",
        "C:a.csv",
        "a\x00.csv",
        "a\n.csv",
        "%2fsecret.csv",
        "\uff0fsecret.csv",
    ],
)
def test_unsafe_filename(name):
    with pytest.raises(IngestionError, match="plain filename"):
        validate_filename(name)


def test_normalizes_filename():
    assert validate_filename(" sales.csv ") == "sales.csv"


def test_size_boundary_precedes_parsing(tmp_path):
    with (tmp_path / "large").open("w+b") as stream:
        stream.truncate(MAX_BYTES + 1)
        with pytest.raises(IngestionError) as error:
            StructuredFileParser().parse(stream, "large.csv", "text/csv")
        assert error.value.code == "file_too_large"
        assert error.value.status == 413
        stream.truncate(MAX_BYTES)
        with pytest.raises(IngestionError) as error:
            StructuredFileParser().parse(stream, "large.csv", "text/csv")
        assert error.value.code != "file_too_large"


def test_zip_expansion_bomb():
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("huge.xml", b"a" * 1024 * 1024)
    with pytest.raises(IngestionError) as error:
        StructuredFileParser().parse(stream, "bad.xlsx", XLSX)
    assert error.value.code == "unsafe_workbook"


def test_xlsx_external_links_and_xml_entities():
    for name, xml, code in [
        ("xl/externalLinks/link.xml", b"<link/>", "unsafe_workbook"),
        (
            "bad.xml",
            b'<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///etc/passwd">]><x>&secret;</x>',
            "malformed_excel",
        ),
    ]:
        stream = BytesIO()
        with ZipFile(stream, "w") as archive:
            archive.writestr(name, xml)
        with pytest.raises(IngestionError) as error:
            StructuredFileParser().parse(stream, "bad.xlsx", XLSX)
        assert error.value.code == code


def test_forged_excel_dimensions_cannot_hide_data():
    stream = BytesIO()
    with ZipFile(BytesIO(workbook_bytes())) as original, ZipFile(stream, "w") as changed:
        for name in original.namelist():
            data = original.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = data.replace(b'ref="A1:B2"', b'ref="A1:A1"')
            changed.writestr(name, data)
    result = StructuredFileParser().parse(stream, "sales.xlsx", XLSX)
    assert (result.row_count, result.column_count) == (1, 2)


def test_missing_excel_values_remain_valid_structure():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["item", "amount"])
    sheet.append(["widget", None])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    assert StructuredFileParser().parse(stream, "sales.xlsx", XLSX).row_count == 1


def test_markup_is_not_accepted_as_csv():
    with pytest.raises(IngestionError) as error:
        StructuredFileParser().parse(BytesIO(b"<html>\n</html>"), "page.csv", "text/csv")
    assert error.value.code == "content_mismatch"


def test_exact_twenty_mib_valid_csv_is_accepted(tmp_path):
    with (tmp_path / "boundary.csv").open("w+b") as stream:
        stream.write(b"text\n")
        remaining = MAX_BYTES - stream.tell()
        while remaining:
            width = min(60_000, remaining)
            stream.write(b"x" * (width - 1) + b"\n")
            remaining -= width
        assert stream.tell() == MAX_BYTES
        result = StructuredFileParser().parse(stream, "boundary.csv", "text/csv")
        assert result.column_count == 1
        assert result.row_count > 1
