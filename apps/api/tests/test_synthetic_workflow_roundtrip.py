"""Use case: Checks fictional business domains through actual authorized workflows.

What it does: Compares independent decimal totals with query, dashboard, saved study,
PDF, Excel and CSV evidence, and confirms a teammate cannot export private studies.
"""

import importlib.util
import json
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from execplus.domain.summary_guard import extract_numbers
from test_studies import add_member, prepare, run
from test_understanding import confirm, context

spec = importlib.util.spec_from_file_location(
    "workflow_kit", Path(__file__).resolve().parents[3] / "scripts/create_synthetic_workflow_kit.py"
)
kit = importlib.util.module_from_spec(spec)
spec.loader.exec_module(kit)


def workbook_evidence(content):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    assert {"Summary", "Methodology", "Evidence", "Read me"} <= set(workbook.sheetnames)
    assert all(sheet.sheet_state == "visible" for sheet in workbook)
    metadata = json.loads("".join(str(row[1]) for row in list(workbook["Evidence"].values)[1:]))
    workbook.close()
    return metadata


@pytest.mark.parametrize(
    "filename,metric,expected,filters",
    [
        ("finance-daily.csv", "revenue", "12714", []),
        ("sales-orders.csv", "revenue", "2340", []),
        ("hr-workforce.csv", "salary", "30900", []),
        (
            "inventory-snapshots.csv",
            "on_hand",
            "329",
            [{"column": "date", "operator": "eq", "value": "2025-01-30"}],
        ),
    ],
    ids=["finance", "sales", "hr", "inventory-snapshot"],
)
def test_business_fixture_calculations_agree_across_authorized_outputs(
    integration, filename, metric, expected, filters
):
    env = integration
    artifacts = kit.artifacts()
    manifest = json.loads(artifacts["manifest.json"])
    entry = next(row for row in manifest["datasets"] if row["file"] == filename)
    original_bytes = artifacts[filename]
    if filename == "finance-daily.csv":
        # Explicit dirty synthetic variant; amounts/dates and their independent oracle remain.
        original_bytes = original_bytes.replace(b",Services,", b", Services ,")
    owner, _, wid, _, _, root = prepare(env, original_bytes, entry["domain"])
    if filename == "finance-daily.csv":
        original = env.client.get(root + "/profile", headers=owner).json()
        body = {"expected_revision_id": original["id"], "trim": True}
        preview = env.client.post(root + "/cleaning/preview", headers=owner, json=body)
        assert preview.status_code == 200, preview.text
        assert preview.json()["revision"]["profile"]["row_count"] == 120
        assert preview.json()["removed_rows"] == 0
        assert env.client.get(root + "/profile", headers=owner).json() == original
        applied = env.client.post(root + "/cleaning/apply", headers=owner, json=body)
        assert applied.status_code == 201, applied.text
        active = applied.json()["revision"]
        assert active["output_checksum"] == preview.json()["revision"]["output_checksum"]
        assert active["output_checksum"] != original["output_checksum"]
        assert active["source_checksum"] == original["source_checksum"]
        assert active["parent_id"] == original["id"]
    current = context(env, owner, root)
    definition = current["definition"]
    for column in definition["columns"]:
        if column["name"] in entry["units"]:
            column["unit"] = entry["units"][column["name"]]
            column["unit_column"] = ""
    definition["metrics"] = [
        {
            "name": "Explicit fixture total",
            "column": metric,
            "aggregation": "sum",
            "filters": filters,
        }
    ]
    assert confirm(env, owner, root, current).status_code == 201
    query = env.client.post(
        root + "/query",
        headers=owner,
        json={
            "metric": metric,
            "aggregation": "sum",
            "group_by": [],
            "filters": filters,
        },
    )
    assert query.status_code == 200, query.text
    result = query.json()
    value = result["rows"][0][0]
    assert Decimal(value) == Decimal(expected)
    dashboard = env.client.post(root + "/dashboard", headers=owner, json={"filters": filters})
    assert dashboard.status_code == 200, dashboard.text
    cards = {card["metric"]: card for card in dashboard.json()["cards"]}
    assert Decimal(cards[metric]["rows"][0][0]) == Decimal(expected)
    saved = run(
        env,
        owner,
        root,
        {
            "kind": "metric",
            "column": metric,
            "aggregation": "sum",
            "filters": filters,
        },
    )
    assert saved.status_code == 201, saved.text
    path = f"/workspaces/{wid}/study-versions/{saved.json()['version']['id']}"
    before = env.client.get(path, headers=owner).json()
    assert before["display"]["rows"] == [[value]]
    xlsx = env.client.get(path + "/export.xlsx", headers=owner)
    assert xlsx.status_code == 200, xlsx.text
    workbook = workbook_evidence(xlsx.content)
    bundle = env.client.get(path + "/export.csv", headers=owner)
    assert bundle.status_code == 200, bundle.text
    with ZipFile(BytesIO(bundle.content)) as archive:
        evidence = json.loads(archive.read("evidence.json"))
    assert workbook["tables"] == evidence["tables"]
    assert any(
        cell["type"] in {"decimal", "integer", "float"}
        and Decimal(cell["value"]) == Decimal(expected)
        for table in evidence["tables"]
        for row in table["rows"]
        for cell in row
    )
    pdf = env.client.get(path + "/export.pdf", headers=owner)
    assert pdf.status_code == 200, pdf.text
    pdf_text = " ".join(page.extract_text() for page in PdfReader(BytesIO(pdf.content)).pages)
    assert Decimal(expected) in extract_numbers(pdf_text) and metric in pdf_text
    assert "evidence" in pdf_text.lower()
    assert env.client.get(path, headers=owner).json() == before
    assert context(env, owner, root)["definition"] == definition
    assert env.client.get(root + "/content", headers=owner).content == original_bytes
    member, _ = add_member(env, owner, wid)
    for format in ("pdf", "xlsx", "csv"):
        assert env.client.get(path + "/export." + format, headers=member).status_code == 404
