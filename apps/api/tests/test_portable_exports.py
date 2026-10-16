"""Use case: Independently opens portable result files and checks fidelity.

What it does: Verifies exact spreadsheets/CSV, formula defenses, image evidence,
real graph figures, output bounds and signed/scatter geometry without web state.
"""

import csv
import json
from dataclasses import replace
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO, StringIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from execplus.domain.exports import ExportChart, ExportContent, ExportSeries, ExportTable
from execplus.domain.ingestion import IngestionError
from execplus.domain.pdf_exports import PdfDocument, PdfSection
from execplus.infrastructure.exports.portable import PortableRenderer

NOW = datetime(2026, 10, 16, tzinfo=timezone.utc)


def content(rows=None, charts=()):
    rows = (
        rows
        if rows is not None
        else (
            (
                Decimal("9007199254740993.000000000000000001"),
                None,
                "",
                False,
                date(2026, 10, 16),
                '=HYPERLINK("https://invalid.test")',
                "\t+SUM(1,2)",
                Decimal("-0.125"),
            ),
        )
    )
    table = ExportTable(
        "Original result", tuple(f"Field {index}" for index in range(len(rows[0]))), rows
    )
    return ExportContent(
        PdfDocument(
            "Fictional exact result",
            NOW,
            (PdfSection("Evidence", ("Source checksum: fictional-source; no new calculation.",)),),
        ),
        (table,),
        charts,
    )


def test_excel_exact_text_cells_and_recoverable_types_without_formula_xml():
    value = PortableRenderer().render(content(), "xlsx")
    workbook = load_workbook(BytesIO(value), data_only=False)
    row = list(workbook["Result 1"].rows)[1]
    assert row[0].value == "9007199254740993.000000000000000001"
    assert row[3].value == "false" and row[4].value == "2026-10-16"
    assert row[5].value.startswith("=HYPERLINK") and row[5].data_type == "s"
    assert row[6].value == "\t+SUM(1,2)" and row[7].value == "-0.125"
    evidence = json.loads("".join(row[1] for row in list(workbook["Evidence"].values)[1:]))
    cells = evidence["tables"][0]["rows"][0]
    assert cells[1] == {"type": "null", "value": ""}
    assert cells[2] == {"type": "string", "value": ""}
    assert cells[0]["type"] == "decimal" and cells[4]["type"] == "date"
    with ZipFile(BytesIO(value)) as archive:
        assert not any("externalLink" in path for path in archive.namelist())
        assert all(
            b"<f>" not in archive.read(path) for path in archive.namelist() if path.endswith(".xml")
        )


def test_csv_bundle_has_every_row_quoted_formula_defenses_and_original_evidence():
    data = content(
        tuple((f"@{index}", Decimal(f"{index}.1234567890123456789")) for index in range(1000))
    )
    value = PortableRenderer().render(data, "csv")
    with ZipFile(BytesIO(value)) as archive:
        rows = list(csv.reader(StringIO(archive.read("result-1.csv").decode("utf-8-sig"))))
        evidence = json.loads(archive.read("evidence.json"))
    assert len(rows) == 1001 and rows[-1] == ["'@999", "999.1234567890123456789"]
    assert evidence["tables"][0]["rows"][-1][0]["value"] == "@999"
    assert "Source checksum" in evidence["sections"][0]["paragraphs"][0]


@pytest.mark.parametrize("unsafe", ["=1+1", " +cmd", "\t@sum(1)", "\r-1+2", "\ufeff=1"])
def test_csv_escapes_formula_looking_strings_and_preserves_them_in_sidecar(unsafe):
    value = PortableRenderer().render(content(((unsafe, Decimal("-0.25")),)), "csv")
    with ZipFile(BytesIO(value)) as archive:
        rows = list(csv.reader(StringIO(archive.read("result-1.csv").decode("utf-8-sig"))))
        evidence = json.loads(archive.read("evidence.json"))
    assert rows[1] == ["'" + unsafe, "-0.25"]
    assert evidence["tables"][0]["rows"][0][0]["value"] == unsafe


@pytest.mark.parametrize("value", ["x" * 32768, "bad\x00text", "😀" * 16384])
def test_excel_refuses_silent_text_truncation_and_illegal_controls(value):
    with pytest.raises(IngestionError) as failure:
        PortableRenderer().render(content(((value,),)), "xlsx")
    assert failure.value.code == "export_text"
    assert PortableRenderer().render(content(((value,),)), "csv")


def test_png_records_exact_all_positions_nulls_and_signed_baseline():
    chart = ExportChart(
        "bar",
        "Signed recorded values",
        "Category",
        "USD",
        ("Negative", "Zero", "Missing", "Positive"),
        (ExportSeries("Value", ("-9007199254740993.1", "0", None, "1")),),
    )
    image = Image.open(BytesIO(PortableRenderer().render(content(charts=(chart,)), "png")))
    assert image.format == "PNG" and image.width == 1600
    recorded = json.loads(image.info["ExecPlus chart"])
    assert recorded["series"][0]["values"] == ["-9007199254740993.1", "0", None, "1"]
    assert (
        json.loads(image.info["ExecPlus evidence"])["tables"][0]["rows"][0][0]["value"]
        == "9007199254740993.000000000000000001"
    )
    assert image.getpixel((1120, 277)) == (23, 108, 100)


def test_scatter_uses_actual_numeric_x_and_line_preserves_null_gap():
    chart = ExportChart(
        "scatter",
        "Uneven x values",
        "X",
        "Y",
        ("0", "1", "100"),
        (ExportSeries("Y", ("1", "1", "1")),),
    )
    image = Image.open(BytesIO(PortableRenderer().render(content(charts=(chart,)), "png")))
    assert image.getpixel((312, 210)) == (23, 108, 100)
    assert image.getpixel((800, 210)) != (23, 108, 100)
    line = replace(chart, kind="line", series=(ExportSeries("Y", ("1", None, "1")),))
    image = Image.open(BytesIO(PortableRenderer().render(content(charts=(line,)), "png")))
    assert image.getpixel((800, 210)) != (23, 108, 100)


def test_diagram_pdf_contains_a_real_figure_and_all_field_text():
    nodes = (
        {"id": "dataset", "kind": "dataset", "label": "This file"},
        {"id": "metrics", "kind": "group", "label": "Measures"},
        *(
            {"id": f"field-{index}", "kind": "field", "label": f"Metric {index}"}
            for index in range(23)
        ),
    )
    edges = tuple({"source": "metrics", "target": f"field-{index}"} for index in range(23))
    chart = ExportChart("diagram", "Data membership", nodes=nodes, edges=edges)
    report = content(charts=(chart,))
    document = replace(
        report.document,
        sections=(PdfSection("All fields", tuple(f"Metric {index}" for index in range(23))),),
    )
    reader = PdfReader(
        BytesIO(PortableRenderer().render(replace(report, document=document), "pdf"))
    )
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "Metric 22" in text and "fields 21-23 of 23" in text
    assert sum(len(page.images) for page in reader.pages) == 3


def test_png_large_result_refuses_without_dropping_rows():
    chart = ExportChart(
        "bar",
        "Many categories",
        labels=tuple(str(index) for index in range(1000)),
        series=(ExportSeries("Value", ("1",) * 1000),),
    )
    with pytest.raises(IngestionError) as failure:
        PortableRenderer().render(content(charts=(chart,)), "png")
    assert failure.value.code == "export_limit"
