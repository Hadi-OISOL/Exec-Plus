"""Use case: Verifies portable exports against real authorized saved results.

What it does: Opens every file format independently, exercises private denial and
revocation races, and rejects browser-supplied data or stale source evidence.
"""

from io import BytesIO
from uuid import UUID
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader
from sqlalchemy import delete, select, update

from execplus.infrastructure.persistence import schema as s
from test_dataset_guidance import ask
from test_pdf_exports import setup
from test_studies import add_member, run
from test_understanding import confirm, context


def download_text(response, format):
    assert response.status_code == 200, response.text
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    if format == "xlsx":
        workbook = load_workbook(BytesIO(response.content))
        return " ".join(
            str(value)
            for sheet in workbook
            if sheet.title.startswith("Result ")
            for row in sheet.values
            for value in row
        )
    if format == "csv":
        with ZipFile(BytesIO(response.content)) as archive:
            return archive.read("evidence.json").decode()
    if format == "png":
        image = Image.open(BytesIO(response.content))
        return image.info["ExecPlus evidence"]
    return " ".join(page.extract_text() for page in PdfReader(BytesIO(response.content)).pages)


@pytest.mark.parametrize("format", ["xlsx", "csv", "png"])
def test_private_thread_exports_exact_rows_with_receipts_and_no_model_calls(integration, format):
    env = integration
    owner, _, wid, root, _, _, url, model = setup(env, grouped=True)
    before = context(env, owner, root)
    response = env.client.get(url.replace(".pdf", "." + format), headers=owner)
    text = download_text(response, format)
    assert "0.100000000000" in text and "0.200000000000" in text
    assert "Karachi" in text and "Lahore" in text
    assert context(env, owner, root) == before and len(model.requests) == 1
    member, _ = add_member(env, owner, wid, "admin")
    assert env.client.get(url.replace(".pdf", "." + format), headers=member).status_code in {
        403,
        404,
    }
    assert env.client.get(url.replace(".pdf", "." + format)).status_code == 401
    history = env.client.get(f"/workspaces/{wid}/audit-history", headers=member).json()
    assert not any(item["action"] == "result.exported" for item in history["events"])


@pytest.mark.parametrize("format", ["xlsx", "csv", "png"])
def test_membership_revocation_while_rendering_refuses_export(integration, monkeypatch, format):
    env = integration
    owner, actor, _, _, _, _, url, _ = setup(env, grouped=True)
    renderer = env.runtime.result_exports.renderer
    original = renderer.render

    def revoke(*args):
        value = original(*args)
        with env.engine.begin() as connection:
            connection.execute(delete(s.memberships).where(s.memberships.c.user_id == actor.id))
        return value

    monkeypatch.setattr(renderer, "render", revoke)
    assert env.client.get(url.replace(".pdf", "." + format), headers=owner).status_code == 404
    with env.engine.connect() as connection:
        assert not connection.execute(
            select(s.audit_events).where(s.audit_events.c.action == "result.exported")
        ).all()


def test_source_change_while_rendering_discards_bytes(integration, monkeypatch):
    env = integration
    owner, _, _, root, _, _, url, _ = setup(env, grouped=True)
    renderer = env.runtime.result_exports.renderer
    original = renderer.render
    definition = context(env, owner, root)
    definition["definition"]["description"] = "Changed while export was rendering"

    def change(*args):
        value = original(*args)
        assert confirm(env, owner, root, definition).status_code == 201
        return value

    monkeypatch.setattr(renderer, "render", change)
    response = env.client.get(url.replace(".pdf", ".csv"), headers=owner)
    assert response.status_code == 409 and response.json()["error"]["code"] == "source_changed"


@pytest.mark.parametrize("format", ["pdf", "xlsx", "csv", "png"])
def test_data_map_export_contains_actual_graph_and_all_saved_roles(integration, format):
    env = integration
    owner, _, _, root, path, _, _, _ = setup(env)
    answer = ask(env, owner, path, "Draw a diagram of this data")
    response = env.client.get(
        path + f"/turns/{answer['turn']['id']}/export.{format}", headers=owner
    )
    text = download_text(response, format)
    assert "amount" in text and "city" in text and "cost" in text
    if format == "pdf":
        assert any(page.images for page in PdfReader(BytesIO(response.content)).pages)
    if format in {"csv", "png"}:
        assert "dataset-structure-v2" in text
    assert context(env, owner, root)["state"] == "confirmed"


@pytest.mark.parametrize("format", ["pdf", "xlsx", "csv", "png"])
def test_current_dashboard_exports_only_replayed_ids_and_rejects_client_values(integration, format):
    env = integration
    owner, _, wid, _, _, answer, _, model = setup(env, grouped=True)
    body = {
        "query_ids": [answer["answer"]["lineage"]["query_id"]],
        "format": format,
        "name": "Current view",
    }
    response = env.client.post(f"/workspaces/{wid}/exports/dashboard", headers=owner, json=body)
    text = download_text(response, format)
    assert "0.100000000000" in text and "0.200000000000" in text
    assert len(model.requests) == 1
    assert (
        env.client.post(
            f"/workspaces/{wid}/exports/dashboard",
            headers=owner,
            json={**body, "rows": [["fake", 999]]},
        ).status_code
        == 422
    )
    assert (
        env.client.post(
            f"/workspaces/{wid}/exports/dashboard",
            headers=owner,
            json={**body, "query_ids": body["query_ids"] * 2},
        ).status_code
        == 422
    )


def test_shared_study_and_board_authorization_and_unshare_render_race(integration, monkeypatch):
    env = integration
    owner, _, wid, root, _, _, _, _ = setup(env)
    created = run(
        env,
        owner,
        root,
        {"kind": "metric", "column": "amount", "aggregation": "sum", "group_by": ["city"]},
    ).json()
    study_id, version_id = created["study"]["id"], created["version"]["id"]
    member, _ = add_member(env, owner, wid)
    url = f"/workspaces/{wid}/study-versions/{version_id}/export.xlsx"
    assert env.client.get(url, headers=member).status_code in {403, 404}
    assert (
        env.client.patch(
            f"/workspaces/{wid}/studies/{study_id}", headers=owner, json={"shared": True}
        ).status_code
        == 204
    )
    assert "0.100000000000" in download_text(env.client.get(url, headers=member), "xlsx")
    board = env.client.post(
        f"/workspaces/{wid}/study-boards",
        headers=owner,
        json={"name": "Shared board", "shared": True, "pins": [version_id]},
    ).json()
    board_url = f"/workspaces/{wid}/study-boards/{board['id']}/export.csv"
    assert "0.200000000000" in download_text(env.client.get(board_url, headers=member), "csv")
    renderer = env.runtime.result_exports.renderer
    original = renderer.render

    def unshare(*args):
        value = original(*args)
        with env.engine.begin() as connection:
            connection.execute(
                update(s.studies).where(s.studies.c.id == UUID(study_id)).values(shared=False)
            )
        return value

    monkeypatch.setattr(renderer, "render", unshare)
    assert env.client.get(board_url, headers=member).status_code in {403, 404}


@pytest.mark.parametrize("kind", ["scenario", "anomaly", "allocation"])
@pytest.mark.parametrize("format", ["pdf", "xlsx", "csv", "png"])
def test_saved_decision_reports_replay_exact_results_and_stay_private(integration, kind, format):
    from test_decision_analysis_integration import create as decision
    from test_scenarios_integration import create as scenario
    from test_studies import prepare

    env = integration
    data = (
        b"city,amount,cost\nKarachi,10,0.05\nLahore,100,0.10\n"
        if kind == "allocation"
        else b"city,amount,cost\n" + b"Karachi,10,0.05\n" * 8 + b"Lahore,100,0.10\n"
    )
    headers, _, wid, _, _, root = prepare(env, data)
    response = (
        scenario(env, headers, root)
        if kind == "scenario"
        else decision(env, headers, root, "anomalies" if kind == "anomaly" else "allocations")
    )
    assert response.status_code == 201, response.text
    saved = response.json()
    url = root.split("/uploads/")[0] + f"/analyses/{saved['run']['id']}/export.{format}?kind={kind}"
    text = download_text(env.client.get(url, headers=headers), format)
    assert "100" in text if kind == "anomaly" else "0.05" in text or "180" in text or "200" in text
    member, _ = add_member(env, headers, wid, "admin")
    assert env.client.get(url, headers=member).status_code in {403, 404}
    reopened = env.client.get(
        root.split("/uploads/")[0] + "/decision-analysis/" + saved["run"]["id"], headers=headers
    )
    assert reopened.json() == saved


@pytest.mark.parametrize("format", ["pdf", "xlsx", "csv", "png"])
def test_measured_pattern_exports_have_recorded_statistics_and_no_causal_claim(integration, format):
    from test_pattern_conversations import setup as pattern_setup

    env = integration
    owner, wid, _, path = pattern_setup(env)
    answer = ask(env, owner, path, "Find patterns in this data")
    assert answer["answer"]["kind"] == "patterns", answer
    url = path + f"/turns/{answer['turn']['id']}/export.{format}"
    text = download_text(env.client.get(url, headers=owner), format)
    assert "1.000000" in text and "40" in text
    if format != "xlsx":
        assert "not causal" in text or "not causality" in text or "not causes" in text
    member, _ = add_member(env, owner, wid)
    assert env.client.get(url, headers=member).status_code in {403, 404}


@pytest.mark.parametrize("format", ["pdf", "xlsx", "csv", "png"])
def test_restricted_export_omits_hidden_field_meanings_and_out_of_scope_rows(integration, format):
    from test_conversational_explorer import ScriptedModel

    env = integration
    owner, actor, wid, root, _, original, old_url, _ = setup(env, grouped=True)
    original_source = original["answer"]["lineage"]["receipt"]["sources"][0]
    meaning = context(env, owner, root)
    for column in meaning["definition"]["columns"]:
        if column["name"] == "cost":
            column["meaning"] = "Secret payroll note never available to the limited user"
    assert confirm(env, owner, root, meaning).status_code == 201
    member, member_actor = add_member(env, owner, wid)
    policy = env.client.put(
        root.split("/uploads/")[0] + "/access-policy",
        headers=owner,
        json={
            "expected_version": 0,
            "enabled": True,
            "grants": [
                {"user_id": str(actor.id), "columns": None, "conditions": []},
                {
                    "user_id": str(member_actor.id),
                    "columns": ["city", "amount"],
                    "conditions": [{"column": "city", "operator": "eq", "value": "Karachi"}],
                },
            ],
        },
    )
    assert policy.status_code == 200, policy.text
    tid = env.client.post(root + "/threads", headers=member).json()["id"]
    path = f"/workspaces/{wid}/threads/{tid}"
    env.runtime.intent_router.model = ScriptedModel(
        {
            "kind": "numerical",
            "plan": {"metric": "amount", "aggregation": "sum", "group_by": ["city"]},
        }
    )
    answer = ask(env, member, path, "Total amount by city?")
    assert answer["answer"] is not None, answer
    url = path + f"/turns/{answer['turn']['id']}/export.{format}"
    response = env.client.get(url, headers=member)
    text = download_text(response, format)
    if format == "xlsx":
        workbook = load_workbook(BytesIO(response.content))
        text = " ".join(str(value) for sheet in workbook for row in sheet.values for value in row)
    assert "Secret payroll note" not in text
    assert "withheld_for_restricted_access" in text
    for key in ("source_checksum", "output_checksum"):
        assert original_source[key][:24] not in text
    assert "Lahore" not in text
    assert "0.100000000000" in text and "0.200000000000" not in text
    assert env.client.get(old_url.replace(".pdf", "." + format), headers=member).status_code in {
        403,
        404,
    }
