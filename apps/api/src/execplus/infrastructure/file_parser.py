"""Use case: Validates bounded single-table CSV and XLSX uploads.

What it does: Rejects unsafe names, formulas, malformed tables, and unsafe workbook structures.
"""

import csv
import re
import unicodedata
from datetime import date, datetime
from io import TextIOWrapper
from pathlib import PurePosixPath
from typing import BinaryIO
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

from execplus.domain.ingestion import FileStructure, IngestionError
from execplus.domain.profiling import TableData

MAX_BYTES = 20 * 1024 * 1024
MAX_CELLS = 1_000_000
MAX_COLUMNS = 1000
MAX_ROWS = 100_000
MIME_TYPES = {
    ".csv": {"text/csv", "application/csv", "text/plain"},
    ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
}


def validate_filename(filename: str) -> str:
    name = unicodedata.normalize("NFKC", filename).strip()
    if (
        not name
        or len(name) > 200
        or name.startswith(".")
        or any(char in name for char in "/\\:%")
        or any(unicodedata.category(char).startswith("C") for char in name)
    ):
        raise IngestionError(
            "unsafe_filename", "Use a plain filename without paths or control characters.", 422
        )
    return name


def validate_cell(value: object) -> None:
    if not isinstance(value, str):
        return
    if any(ord(char) < 32 and char not in "\t\r\n" for char in value):
        raise IngestionError(
            "unsafe_content", "The file contains unsupported control characters.", 422
        )
    stripped = value.lstrip()
    if stripped.startswith(("=", "+", "@")) or (
        stripped.startswith("-") and not re.fullmatch(r"-\d+(\.\d+)?([eE][+-]?\d+)?", stripped)
    ):
        raise IngestionError(
            "unsafe_formula", "Formula-like cells are unsupported. Upload values only.", 422
        )


def validate_header(values: list[object]) -> None:
    headers = [str(value).strip().casefold() if value is not None else "" for value in values]
    if not headers or any(not header for header in headers) or len(set(headers)) != len(headers):
        raise IngestionError(
            "invalid_header", "Provide one nonempty, unique header for each column.", 422
        )


def validate_dimensions(rows: int, columns: int) -> None:
    if columns > MAX_COLUMNS or rows > MAX_ROWS or rows * columns > MAX_CELLS:
        raise IngestionError(
            "unsupported_structure", "The table exceeds supported row, column, or cell limits.", 422
        )


class StructuredFileParser:
    def parse(self, content: BinaryIO, filename: str, content_type: str) -> FileStructure:
        filename = validate_filename(filename)
        extension = PurePosixPath(filename).suffix.lower()
        if (
            extension not in MIME_TYPES
            or content_type.split(";")[0].lower() not in MIME_TYPES[extension]
        ):
            raise IngestionError(
                "unsupported_format",
                "Choose a CSV or XLSX file with its matching content type.",
                422,
            )
        content.seek(0, 2)
        size = content.tell()
        content.seek(0)
        if size > MAX_BYTES:
            raise IngestionError("file_too_large", "Files must be at most 20 MiB.", 413)
        if size == 0:
            raise IngestionError("empty_file", "The file is empty.", 422)
        prefix = content.read(8)
        content.seek(0)
        if extension == ".csv":
            if prefix.startswith((b"PK", b"\xd0\xcf\x11\xe0", b"MZ", b"%PDF")):
                raise IngestionError("content_mismatch", "The file content is not CSV text.", 422)
            preview = content.read(4096).lstrip().lower()
            content.seek(0)
            if preview.startswith((b"<html", b"<!doctype", b"<?xml", b"<script")):
                raise IngestionError(
                    "content_mismatch", "Markup documents are not supported CSV files.", 422
                )
            while chunk := content.read(65536):
                if b"\x00" in chunk:
                    raise IngestionError("unsafe_content", "Null bytes are unsupported.", 422)
            content.seek(0)
            return self._csv(content)
        if not prefix.startswith(b"PK\x03\x04"):
            raise IngestionError(
                "malformed_excel", "The file is not a readable XLSX workbook.", 422
            )
        return self._xlsx(content)

    def read_table(self, content: BinaryIO, format: str) -> TableData:
        content.seek(0)
        if format == "csv":
            wrapper = TextIOWrapper(content, encoding="utf-8-sig", newline="")
            try:
                reader = csv.reader(wrapper, strict=True)
                return TableData(tuple(next(reader)), tuple(tuple(row) for row in reader))
            finally:
                wrapper.detach()
                content.seek(0)
        workbook = load_workbook(content, read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()
            iterator = sheet.iter_rows(values_only=True)
            headers = tuple(str(value) for value in next(iterator))

            def text(value: object) -> str:
                if value is None:
                    return ""
                if isinstance(value, datetime):
                    return (
                        value.date().isoformat()
                        if value.time().isoformat() == "00:00:00"
                        else value.isoformat()
                    )
                if isinstance(value, date):
                    return value.isoformat()
                return str(value)

            rows = tuple(
                tuple(text(row[i]) if i < len(row) else "" for i in range(len(headers)))
                for row in iterator
            )
            return TableData(headers, rows)
        finally:
            workbook.close()
            content.seek(0)

    def _csv(self, content: BinaryIO) -> FileStructure:
        wrapper = TextIOWrapper(content, encoding="utf-8-sig", newline="")
        try:
            reader = csv.reader(wrapper, strict=True)
            header = next(reader, [])
            validate_header(list(header))
            columns = len(header)
            validate_dimensions(1, columns)
            for cell in header:
                validate_cell(cell)
            rows = 0
            for row in reader:
                if len(row) != columns or not any(cell.strip() for cell in row):
                    raise IngestionError(
                        "malformed_csv",
                        "CSV rows must have consistent columns and cannot be empty.",
                        422,
                    )
                rows += 1
                validate_dimensions(rows + 1, columns)
                for cell in row:
                    validate_cell(cell)
            if rows == 0:
                raise IngestionError(
                    "empty_table", "Include at least one data row below the header.", 422
                )
            return FileStructure("csv", rows, columns, 0)
        except (csv.Error, UnicodeError) as error:
            raise IngestionError(
                "malformed_csv", "Use a well-formed UTF-8 CSV with consistent quoting.", 422
            ) from error
        finally:
            wrapper.detach()
            content.seek(0)

    def _inspect_archive(self, archive: ZipFile) -> None:
        files = archive.infolist()
        names = [file.filename for file in files]
        if len(files) > 512 or len(set(names)) != len(names):
            raise IngestionError(
                "unsafe_workbook", "The workbook archive has unsupported entries.", 422
            )
        if sum(file.file_size for file in files) > 100 * 1024 * 1024:
            raise IngestionError("unsafe_workbook", "The expanded workbook is too large.", 422)
        for file in files:
            if (
                file.flag_bits & 1
                or file.file_size > 30 * 1024 * 1024
                or file.file_size > max(file.compress_size, 1) * 200
                or file.filename.startswith("/")
                or ".." in PurePosixPath(file.filename).parts
                or "\\" in file.filename
            ):
                raise IngestionError(
                    "unsafe_workbook",
                    "Encrypted, unsafe, or highly compressed workbooks are unsupported.",
                    422,
                )
            lower = file.filename.lower()
            if any(
                part in lower for part in ("vbaproject", "externallinks", "embeddings", "activex")
            ):
                raise IngestionError(
                    "unsafe_workbook",
                    "Macros, embedded objects, and external links are unsupported.",
                    422,
                )
            if lower.endswith((".xml", ".rels")):
                cells = 0
                with archive.open(file) as stream:
                    for _, element in iterparse(stream, events=("end",), forbid_dtd=True):
                        tag = element.tag.rsplit("}", 1)[-1]
                        if tag == "mergeCell":
                            raise IngestionError(
                                "merged_cells", "Remove merged cells before uploading.", 422
                            )
                        if tag == "f":
                            raise IngestionError(
                                "unsafe_formula",
                                "Excel formulas are unsupported. Upload values only.",
                                422,
                            )
                        if element.attrib.get("TargetMode") == "External":
                            raise IngestionError(
                                "unsafe_workbook",
                                "External workbook relationships are unsupported.",
                                422,
                            )
                        if tag == "c":
                            if "r" in element.attrib:
                                row, column = coordinate_to_tuple(element.attrib["r"])
                                validate_dimensions(row, column)
                            cells += 1
                            if cells > MAX_CELLS:
                                raise IngestionError(
                                    "unsupported_structure",
                                    "The workbook contains too many cells.",
                                    422,
                                )
                        element.clear()

    def _xlsx(self, content: BinaryIO) -> FileStructure:
        try:
            with ZipFile(content) as archive:
                self._inspect_archive(archive)
            content.seek(0)
            workbook = load_workbook(content, read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.sheetnames) != 1 or len(workbook.worksheets) != 1:
                    raise IngestionError(
                        "multiple_sheets",
                        "Upload a workbook containing exactly one worksheet.",
                        422,
                    )
                sheet = workbook.worksheets[0]
                sheet.reset_dimensions()
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, ())
                validate_header(list(header))
                columns = len(header)
                for cell in header:
                    validate_cell(cell)
                count = 0
                for row in rows:
                    if len(row) > columns:
                        raise IngestionError(
                            "unsupported_structure",
                            "Every Excel row must match the header width.",
                            422,
                        )
                    if not any(cell is not None for cell in row):
                        raise IngestionError(
                            "unsupported_structure",
                            "Empty rows inside the table are unsupported.",
                            422,
                        )
                    count += 1
                    validate_dimensions(count + 1, columns)
                    for cell in row:
                        validate_cell(cell)
                if count == 0:
                    raise IngestionError(
                        "empty_table", "Include at least one data row below the header.", 422
                    )
                return FileStructure("xlsx", count, columns, 1)
            finally:
                workbook.close()
        except IngestionError:
            raise
        except (
            BadZipFile,
            DefusedXmlException,
            ValueError,
            KeyError,
            OSError,
            TypeError,
            SyntaxError,
        ) as error:
            raise IngestionError(
                "malformed_excel", "The file is not a safe, readable XLSX workbook.", 422
            ) from error
        finally:
            content.seek(0)
