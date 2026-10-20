"""Use case: Delivers exact results as spreadsheets, CSV evidence bundles and charts.

What it does: Keeps types and provenance, disables spreadsheet formula interpretation,
refuses truncation and uses only local trusted renderers.
"""

import csv
import json
import re
from contextlib import suppress
from dataclasses import asdict, replace
from datetime import timezone
from io import BytesIO, StringIO
from typing import Any
from zipfile import ZIP_DEFLATED

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import ExcelWriter

from execplus.domain.exports import (
    MAX_EXPORT_BYTES,
    ExportContent,
    cell_type,
    exact_text,
    validate_export,
)
from execplus.domain.ingestion import IngestionError
from execplus.infrastructure.exports.chart_png import png_chart
from execplus.infrastructure.exports.reportlab_pdf import ReportLabPdfRenderer
from execplus.infrastructure.exports.reproducible_zip import ReproducibleZip

FORMULA = re.compile(r"^[\s\ufeff]*[=+@-]")
NUMERIC = re.compile(r"^-?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")


def worksheet_name(value: str, used: set[str]) -> str:
    stem = re.sub(r"[\\/*?:\[\]\x00-\x1f]", " ", value).strip().strip("'").strip() or "Result"

    def clipped(text: str, length: int) -> str:
        return text.encode("utf-16-le")[: length * 2].decode("utf-16-le", errors="ignore")

    candidate = clipped(stem, 31).rstrip("'").strip() or "Result"
    sequence = 2
    while candidate.casefold() in used:
        suffix = f" ({sequence})"
        candidate = clipped(stem, 31 - len(suffix)) + suffix
        sequence += 1
    used.add(candidate.casefold())
    return candidate


def manifest(content: ExportContent) -> dict[str, Any]:
    return {
        "version": "execplus-export-v1",
        "title": content.document.title,
        "generated_at": content.document.generated_at.isoformat(),
        "sections": [asdict(section) for section in content.document.sections],
        "tables": [
            {
                "name": table.name,
                "columns": table.columns,
                "rows": [
                    [{"type": cell_type(value), "value": exact_text(value)} for value in row]
                    for row in table.rows
                ],
            }
            for table in content.tables
        ],
        "charts": [asdict(chart) for chart in content.charts],
        "notice": (
            "Exact scalar strings and original types are retained here. "
            "CSV may prefix formula-like text with an apostrophe; Excel cells are explicitly text. "
            "Null and empty string differ by type. CSV readers may infer numbers or dates; "
            "use this evidence or the text-only Excel workbook for exact originals. "
            "A downloaded copy cannot be recalled."
        ),
    }


class PortableRenderer:
    def __init__(self, pdf: ReportLabPdfRenderer | None = None) -> None:
        self.pdf = pdf or ReportLabPdfRenderer()

    def render(self, content: ExportContent, format: str, chart_index: int = 0) -> bytes:
        validate_export(content)
        if format == "pdf":
            value = self.pdf.render(replace(content.document, charts=content.charts))
        else:
            if format not in {"xlsx", "csv", "png"}:
                raise IngestionError("export_format", "Choose PDF, Excel, CSV or PNG.", 422)
            evidence = manifest(content)
            if format == "xlsx":
                value = self.xlsx(content, evidence)
            elif format == "csv":
                value = self.csv_bundle(content, evidence)
            else:
                if chart_index < 0 or chart_index >= len(content.charts):
                    raise IngestionError(
                        "export_chart", "Choose one of this result's available charts.", 422
                    )
                value = png_chart(content.charts[chart_index], evidence)
        if len(value) > MAX_EXPORT_BYTES:
            raise IngestionError(
                "export_limit", "This export exceeds 16 MiB. Choose a smaller result.", 422
            )
        return value

    def csv_bundle(self, content: ExportContent, evidence: dict[str, Any]) -> bytes:
        output = BytesIO()
        with ReproducibleZip(output, "w", compression=ZIP_DEFLATED) as archive:
            for index, table in enumerate(content.tables, 1):
                stream = StringIO(newline="")
                writer = csv.writer(stream, quoting=csv.QUOTE_ALL, lineterminator="\r\n")
                for row in (table.columns, *table.rows):
                    cells = []
                    for item in row:
                        value = exact_text(item)
                        cells.append(
                            "'" + value
                            if FORMULA.match(value) and not NUMERIC.fullmatch(value)
                            else value
                        )
                    writer.writerow(cells)
                archive.writestr(f"result-{index}.csv", stream.getvalue().encode("utf-8-sig"))
            archive.writestr(
                "evidence.json", json.dumps(evidence, ensure_ascii=False, default=str, indent=2)
            )
            archive.writestr(
                "README.txt",
                evidence["notice"]
                + (
                    "\nEach result-N.csv corresponds to table N in evidence.json. Every "
                    "returned row is included.\n"
                ),
            )
        return output.getvalue()

    def xlsx(self, content: ExportContent, evidence: dict[str, Any]) -> bytes:
        def check(value: str, *, long_paragraph: bool = False) -> None:
            try:
                length = len(value.encode("utf-16-le")) // 2
            except UnicodeEncodeError:
                raise IngestionError(
                    "export_text", "This text contains invalid Unicode.", 422
                ) from None
            if (length > 32767 and not long_paragraph) or re.search(
                r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value
            ):
                raise IngestionError(
                    "export_text",
                    "An Excel cell is too long or contains unsupported control characters. "
                    "Use the CSV evidence bundle.",
                    422,
                )

        # Validate before openpyxl creates any temporary worksheet files.
        check(content.document.title)
        check(evidence["notice"])
        for section in content.document.sections:
            check(section.title)
            for paragraph in section.paragraphs:
                check(paragraph, long_paragraph=True)
        for table in content.tables:
            for row in ((table.name,), table.columns, *table.rows):
                for item in row:
                    check(exact_text(item))
        workbook = Workbook(write_only=True)
        try:
            return self._xlsx_workbook(workbook, content, evidence)
        finally:
            # A failed append/archive write must not retain source XML until process exit.
            # These writer attributes belong to the pinned openpyxl write-only API internals.
            for sheet in workbook.worksheets:
                writer = getattr(sheet, "_writer", None)
                if writer is None:
                    continue
                with suppress(Exception):
                    if not sheet.closed:
                        sheet.close()
                with suppress(Exception):
                    writer.close()
                with suppress(FileNotFoundError):
                    writer.cleanup()
            workbook.close()

    def _xlsx_workbook(
        self, workbook: Any, content: ExportContent, evidence: dict[str, Any]
    ) -> bytes:
        timestamp = content.document.generated_at.astimezone(timezone.utc).replace(tzinfo=None)
        workbook.properties.created = timestamp
        workbook.properties.modified = timestamp
        workbook.properties.creator = "ExecPlus"
        workbook.properties.title = content.document.title
        workbook.properties.description = "Exact authorized results; original types in Evidence."

        def append(sheet: Any, values: tuple[str, ...], *, header: bool = False) -> None:
            cells = []
            for value in values:
                if len(value.encode("utf-16-le")) // 2 > 32767 or re.search(
                    r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value
                ):
                    raise IngestionError(
                        "export_text",
                        (
                            "An Excel cell exceeds its text limit or contains unsupported "
                            "control characters. Use the CSV evidence bundle."
                        ),
                        422,
                    )
                cell = WriteOnlyCell(sheet, value=value)
                cell.data_type = "s"
                cell.number_format = "@"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if header:
                    cell.font = Font(bold=True, color="176C64")
                    cell.fill = PatternFill("solid", fgColor="E9F1ED")
                cells.append(cell)
            sheet.append(cells)

        summary = workbook.create_sheet("Summary")
        summary.column_dimensions["A"].width = 25
        summary.column_dimensions["B"].width = 80
        append(summary, ("Report", content.document.title), header=True)
        append(summary, ("Generated at (UTC)", timestamp.isoformat() + "Z"))
        append(summary, ("Included tables", str(len(content.tables))))
        append(summary, ("Exact values", evidence["notice"]))
        used = {"summary", "methodology", "evidence", "read me", "history"}
        table_sheets = []
        for table in content.tables:
            name = worksheet_name(table.name, used)
            table_sheets.append(name)
            append(summary, (name, table.name, f"{len(table.rows)} rows"))
            sheet = workbook.create_sheet(name)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(table.columns))}{len(table.rows) + 1}"
            )
            for index in range(1, len(table.columns) + 1):
                sheet.column_dimensions[get_column_letter(index)].width = 26
            append(sheet, table.columns, header=True)
            for row in table.rows:
                append(sheet, tuple(exact_text(value) for value in row))
        methodology = workbook.create_sheet("Methodology")
        methodology.column_dimensions["A"].width = 30
        methodology.column_dimensions["B"].width = 90
        methodology.column_dimensions["C"].width = 24
        append(
            methodology,
            ("Section", "Saved methodology, caveats and references", "Paragraph / part"),
            header=True,
        )
        for section in content.document.sections:
            for index, paragraph in enumerate(section.paragraphs, 1):
                # 15,000 Unicode characters use at most 30,000 UTF-16 code units.
                chunks = [
                    paragraph[start : start + 15000] for start in range(0, len(paragraph), 15000)
                ] or [""]
                for part, chunk in enumerate(chunks, 1):
                    append(
                        methodology, (section.title, chunk, f"{index} / {part} of {len(chunks)}")
                    )
        meta = workbook.create_sheet("Evidence")
        append(meta, ("Part", "Exact JSON text (concatenate in order)"), header=True)
        serialized = json.dumps(evidence, ensure_ascii=True, default=str, indent=2)
        for start in range(0, len(serialized), 30000):
            append(meta, (str(start // 30000 + 1), serialized[start : start + 30000]))
        guide = workbook.create_sheet("Read me")
        append(guide, ("ExecPlus exact result export",), header=True)
        append(guide, (evidence["notice"],))
        for name, table in zip(table_sheets, content.tables, strict=True):
            append(guide, (name, table.name, f"{len(table.rows)} rows"))
        output = BytesIO()
        # save_workbook() overwrites modified with the wall clock. The existing
        # writer accepts explicit metadata and our deterministic ZIP container.
        with ReproducibleZip(output, "w", compression=ZIP_DEFLATED) as archive:
            ExcelWriter(workbook, archive).save()
        return output.getvalue()
